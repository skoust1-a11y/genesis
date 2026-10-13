"""Backtest of the NIFTY liquidity sweep signals as an Excel workbook: Summary, Yearly, Trades, Daily.

Data: Yahoo's last 60 days of 5m bars (default; VWAP from NIFTYBEES, as the live alerts), or with --history the
Kaggle NIFTY 1-minute CSV rolled up to 5m and topped up from Groww's 5m chart (late Dec 2025 on). The index has no
volume, so --history uses the session's time-weighted average price in place of VWAP.

Entries: app.liquidity_signals.scan, the code the live alerts run. Exits: app.liquidity_signals.exit_of (stop,
breakeven after +1R or target touched by a 5m high/low, else the 15:10 close). Options: `--lots` lots (fewer when the
trading capital cannot pay the premium) of the OTM_STEPS-out-of-the-money CE/PE to the expiry in force (monthly Thursday before 11 Feb 2019, weekly
Thursday until Aug 2025, weekly Tuesday since), at the lot size of the day, Black-76 at the previous day's India VIX
(IV held constant), today's charges, no slippage. --book-profits keeps the trading capital at --capital: losses come
out of it, a winning trade first refills it, and the rest of the profit moves to a wallet.

    curl -L -o nifty.zip https://www.kaggle.com/api/v1/datasets/download/mavrick136/10-years-of-nifty-50-1-minute-data-20152025
    uv run --with openpyxl python scripts/backtest_liquidity.py [--history data/intraday/NIFTY_1MIN_2015-2025.csv]
        [--years 10] [--capital 100000] [--lots 1] [--book-profits] [--out FILE]
"""

from __future__ import annotations

import argparse
import csv
import sys
from bisect import bisect_left
from collections import defaultdict
from datetime import date, datetime, timedelta
from datetime import time as dtime
from pathlib import Path

import httpx

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.api.market import IST, _chart, add_vwap, parse_yahoo_chart
from app.config import Settings, get_settings
from app.engine.metrics import performance
from app.liquidity_signals import (
    BREAK_DAYS,
    EXIT,
    INDEX,
    LOSS_STREAK,
    MIN_RISK,
    OTM_STEPS,
    R_MULT,
    VWAP_PROXY,
    estimated_plan,
    exit_of,
    scan,
)
from app.market.hours import holidays
from app.market.nse import load_lot_sizes, lot_size_on
from app.paper import order_charges
from app.trading.option_pnl import black76

WEEKLY_FROM, TUESDAY_FROM = date(2019, 2, 11), date(2025, 9, 1)
GROWW = (
    "https://groww.in/v1/api/charting_service/v2/chart/exchange/NSE/segment/CASH/NIFTY"
)


# Data


def yahoo_data() -> tuple[list[dict], dict, list[dict]]:
    bars = [
        b
        for b in parse_yahoo_chart(_chart(INDEX, "5m", "60d"))
        if None not in (b["open"], b["high"], b["low"])
    ]
    above = {b["date"]: b["close"] > b["vwap"] if b["vwap"] else None
             for b in add_vwap(parse_yahoo_chart(_chart(VWAP_PROXY, "5m", "60d")))}  # fmt: skip
    return (
        bars,
        above,
        [d for d in parse_yahoo_chart(_chart(INDEX, "1d", "6mo")) if d["high"]],
    )


def kaggle_5m(path: Path) -> list[dict]:
    """5m bars (labelled by start time) from the Kaggle 1-minute CSV; its 15:30 row repeats 15:29 and is dropped.
    The file also fills every NSE holiday with the previous session's last minute repeated all day: those sessions
    are dropped, or the next day's previous-day high/low would be that one minute's range."""
    out: list[dict] = []
    with path.open() as f:
        for row in csv.DictReader(f):
            t = row["datetime"][:16]
            if not "09:15" <= t[11:] < "15:30" or "" in (
                row["open"],
                row["high"],
                row["low"],
                row["close"],
            ):
                continue
            start = f"{t[:14]}{int(t[14:16]) // 5 * 5:02d}"
            o, h, low, c = (float(row[k]) for k in ("open", "high", "low", "close"))
            if out and out[-1]["date"] == start:
                b = out[-1]
                b["high"], b["low"], b["close"] = (
                    max(b["high"], h),
                    min(b["low"], low),
                    c,
                )
            else:
                out.append(
                    {"date": start, "open": o, "high": h, "low": low, "close": c}
                )
    days: dict[str, set] = defaultdict(set)
    for b in out:
        days[b["date"][:10]].add((b["open"], b["high"], b["low"], b["close"]))
    return [b for b in out if len(days[b["date"][:10]]) > 1]


def groww_5m(start: date, end: date) -> list[dict]:
    """NIFTY 5m bars from Groww's public chart service (it keeps about a year), 10 days per request."""
    out, day = [], start
    while day <= end:
        stop = min(day + timedelta(days=10), end + timedelta(days=1))
        ms = [
            int(datetime.combine(d, dtime(9), IST).timestamp() * 1000)
            for d in (day, stop)
        ]
        r = httpx.get(GROWW, params={"startTimeInMillis": ms[0], "endTimeInMillis": ms[1], "intervalInMinutes": 5},
                      headers={"User-Agent": "Mozilla/5.0"}, timeout=20)  # fmt: skip
        r.raise_for_status()
        for ts, o, h, low, c, *_ in r.json()["candles"]:
            out.append({"date": datetime.fromtimestamp(ts, IST).strftime("%Y-%m-%d %H:%M"),
                        "open": o, "high": h, "low": low, "close": c})  # fmt: skip
        day = stop
    return out


def twap_above(bars: list[dict]) -> dict[str, bool]:
    """Whether each 5m bar closed above the session's running mean typical price (VWAP without volume)."""
    out, day, total, n = {}, None, 0.0, 0
    for b in bars:
        if b["date"][:10] != day:
            day, total, n = b["date"][:10], 0.0, 0
        total, n = total + (b["high"] + b["low"] + b["close"]) / 3, n + 1
        out[b["date"]] = b["close"] > total / n
    return out


def history_data(path: Path) -> tuple[list[dict], dict, list[dict]]:
    bars = kaggle_5m(path)
    seen, last = {b["date"] for b in bars}, date.fromisoformat(bars[-1]["date"][:10])
    bars += [
        b
        for b in groww_5m(last + timedelta(days=1), date.today())
        if b["date"] not in seen
    ]
    daily: dict[str, dict] = {}
    for b in bars:
        d = daily.setdefault(
            b["date"][:10], {"date": b["date"][:10], "high": b["high"], "low": b["low"]}
        )
        d["high"], d["low"] = max(d["high"], b["high"]), min(d["low"], b["low"])
    return bars, twap_above(bars), list(daily.values())


# Rules


def expiry_for(day: date, trading: set[str], settings: Settings) -> date:
    """The nearest NIFTY expiry on or after `day`, moved to the trading day before when it falls on a holiday."""
    if day < WEEKLY_FROM:
        exp = day
        for _ in range(
            2
        ):  # this month's last Thursday, or next month's once it has passed
            nxt = date(exp.year + exp.month // 12, exp.month % 12 + 1, 1)
            exp = nxt - timedelta(days=(nxt.weekday() - 3) % 7 or 7)
            if exp >= day:
                break
            exp = nxt
    else:
        thu = day + timedelta(days=(3 - day.weekday()) % 7)
        exp = (
            thu if thu < TUESDAY_FROM else day + timedelta(days=(1 - day.weekday()) % 7)
        )
    off, last = holidays(settings), max(trading)
    while exp > day and (
        exp.weekday() >= 5
        or exp.isoformat() in off
        or exp.isoformat() <= last
        and exp.isoformat() not in trading
    ):
        exp -= timedelta(days=1)  # fmt: skip
    return exp


def option_exit(plan: dict, spot: float, at: str, settings: Settings) -> float:
    strike, kind = plan["option"].split()
    expires = datetime.combine(
        date.fromisoformat(plan["expiry"]),
        dtime.fromisoformat(settings.market_close),
        IST,
    )
    left = expires - datetime.strptime(at, "%Y-%m-%d %H:%M").replace(tzinfo=IST)
    return black76(
        spot,
        float(strike),
        max(left.total_seconds(), 0) / 86400 / 365,
        plan["iv_pct"] / 100,
        kind,
    )


def backtest(settings: Settings, bars: list[dict], above: dict, daily: list[dict], vix: dict[str, float],
             since: str = "", stop_at: str = "swing") -> tuple[list[dict], list[str]]:  # fmt: skip
    """Every signal from `since` (before today) with its index exit and 1-unit option premiums; and the days."""
    by_day: dict[str, list[dict]] = defaultdict(list)
    for b in bars:
        by_day[b["date"][:10]].append(b)
    daily_dates, vix_dates = [d["date"][:10] for d in daily], sorted(vix)
    lots_table = load_lot_sizes(Path(settings.data_dir) / "lot_sizes" / "nse.yaml")
    trading, today = set(by_day), datetime.now(IST).strftime("%Y-%m-%d")
    trades, days = [], []
    for d in sorted(k for k in by_day if since <= k < today):
        i, j = bisect_left(daily_dates, d), bisect_left(vix_dates, d)
        if not i or not j:
            continue
        days.append(d)
        prev, iv = daily[i - 1], vix[vix_dates[j - 1]]
        expiry = expiry_for(date.fromisoformat(d), trading, settings)
        if expiry.isoformat() == d:
            continue  # no trades on expiry day
        for s in scan(by_day[d], above, prev["high"], prev["low"], stop_at=stop_at):
            reason, price, at = exit_of(by_day[d], s)
            points = price - s["entry"] if s["side"] == "Long" else s["entry"] - price
            plan = estimated_plan(s, settings, iv, expiry)
            trades.append({
                "Date": d, "Side": s["side"], "Entry time": s["time"][11:], "Entry": s["entry"], "Stop": s["stop"],
                "Target": s["target"], "Risk pts": s["risk"], "Exit time": at[11:], "Exit reason": reason,
                "Exit": round(price, 2), "Points": round(points, 2), "R": round(points / s["risk"], 2),
                "Option": plan["option"], "Expiry": plan["expiry"], "VIX": plan["iv_pct"],
                "Lot size": lot_size_on(lots_table, "NIFTY", d) or 75,  # 75 from Nov 2015 until the table starts
                "Premium in": plan["entry_premium"], "Premium at stop": plan["exits"][0]["premium"],
                "Premium out": round(option_exit(plan, price, at, settings), 2),
                "Why": " | ".join(s["why"]),
            })  # fmt: skip
    return trades, days


# Money


def ledger(trades: list[dict], days: list[str], settings: Settings, capital: float, lots: int, book: bool,
           risk_pct: float | None = None) -> list[dict]:  # fmt: skip
    """Sizes and settles each trade in entry order. Lots: `risk_pct`% of the current trading capital / the option's
    loss per lot at the stop (charges in) when set, else `lots`; never more than the trading capital can pay. Risking
    a fixed rupee amount instead drained the capital to Rs 3,887 over 2016-2026 on Rs 50,000. LOSS_STREAK losing
    trades in a row (index points) skip the rest of that day and the next BREAK_DAYS trading days. With `book`, a
    winning trade first refills the trading capital to `capital` and the rest moves to the wallet (booking every win
    in full drains the capital: losses would never be won back). Returns one row per day."""
    trades.sort(key=lambda t: (t["Date"], t["Entry time"]))
    cash, wallet, streak, paused = capital, 0.0, 0, ""
    for t in trades:
        kind, lot = t["Option"].split()[1], t["Lot size"]
        unit = t["Premium in"] * lot
        n_cash = max(int(cash // unit), 0) if unit > 0 else lots
        n, skip = min(lots, n_cash), ""
        if risk_pct is not None:
            risk = cash * risk_pct / 100
            loss = (t["Premium in"] - t["Premium at stop"]) * lot + sum(
                order_charges(settings, kind, side, p, lot, t["Date"])["total"]
                for side, p in (("BUY", t["Premium in"]), ("SELL", t["Premium at stop"])))  # fmt: skip
            n = min(int(risk // loss) if loss > 0 else 0, n_cash)
            skip = (
                f"1 lot risks over {risk_pct:g}% of capital"
                if loss <= 0 or risk < loss
                else ""
            )
        if t["Date"] <= paused:
            n, skip = 0, f"{LOSS_STREAK}-loss break"
        elif not n and not skip:
            skip = "capital too low"
        qty, pnl = n * lot, 0.0
        if qty:
            charges = (order_charges(settings, kind, "BUY", t["Premium in"], qty, t["Date"])["total"]
                       + order_charges(settings, kind, "SELL", t["Premium out"], qty, t["Date"])["total"])  # fmt: skip
            pnl = (t["Premium out"] - t["Premium in"]) * qty - charges
            streak = streak + 1 if t["Points"] < 0 else 0
            if streak == LOSS_STREAK:
                k = bisect_left(days, t["Date"]) + BREAK_DAYS
                streak, paused = 0, days[k] if k < len(days) else "9999"
        booked = (
            max(cash + pnl - capital, 0.0) if book and pnl > 0 else 0.0
        )  # refill losses first, book the rest
        cash, wallet = cash + pnl - booked, wallet + booked
        t.update({"Lots": n, "Skipped": skip, "Qty": qty, "Futures Rs": round(t["Points"] * qty),
                  "Option Rs": round(pnl), "Booked Rs": round(booked), "Trading capital Rs": round(cash),
                  "Wallet Rs": round(wallet), "Total Rs": round(cash + wallet)})  # fmt: skip
    rows, k, cash, wallet = [], 0, capital, 0.0
    for d in days:
        today = []
        while k < len(trades) and trades[k]["Date"] == d:
            today.append(trades[k])
            k += 1
        if today:
            cash, wallet = today[-1]["Trading capital Rs"], today[-1]["Wallet Rs"]
        rows.append({"Date": d, "Trades": len(today), "Points": round(sum(t["Points"] for t in today), 2),
                     "Option Rs": sum(t["Option Rs"] for t in today), "Booked Rs": sum(t["Booked Rs"] for t in today),
                     "Trading capital Rs": round(cash), "Wallet Rs": round(wallet), "Total Rs": round(cash + wallet)})  # fmt: skip
    return rows


def yearly(trades: list[dict], days: list[dict]) -> list[dict]:
    out = []
    for year in sorted({d["Date"][:4] for d in days}):
        ts = [t for t in trades if t["Date"][:4] == year and t["Lots"]]
        end = [d for d in days if d["Date"][:4] == year][-1]
        out.append({"Year": year, "Trades": len(ts), "Wins": sum(t["Points"] > 0 for t in ts),
                    "Win rate %": round(sum(t["Points"] > 0 for t in ts) / len(ts) * 100, 1) if ts else None,
                    "Points": round(sum(t["Points"] for t in ts), 2), "R": round(sum(t["R"] for t in ts), 2),
                    "Option Rs": sum(t["Option Rs"] for t in ts), "Booked Rs": sum(t["Booked Rs"] for t in ts),
                    "Trading capital Rs": end["Trading capital Rs"], "Wallet Rs": end["Wallet Rs"],
                    "Total Rs": end["Total Rs"]})  # fmt: skip
    return out


def money(
    trades: list[dict], days: list[dict], capital: float, sizing: str, book: bool
) -> list[tuple]:
    if not days:
        return []
    end, years = (
        days[-1],
        (
            date.fromisoformat(days[-1]["Date"]) - date.fromisoformat(days[0]["Date"])
        ).days
        / 365.25,
    )
    high, low = (
        max(days, key=lambda d: d["Total Rs"]),
        min(days, key=lambda d: d["Total Rs"]),
    )
    low_cash = min(days, key=lambda d: d["Trading capital Rs"])
    cagr = (
        (end["Total Rs"] / capital) ** (1 / years) - 1
        if years > 0 and end["Total Rs"] > 0
        else None
    )
    return [(f"Your money ({sizing}{f', profit above Rs {capital:,.0f} booked to the wallet' if book else ''})", ""),
            ("Starting capital Rs", round(capital)),
            ("Trading capital at the end Rs", end["Trading capital Rs"]),
            ("Wallet (profit booked) Rs", end["Wallet Rs"]),
            ("Total (capital + wallet) Rs", end["Total Rs"]),
            ("Change Rs", end["Total Rs"] - round(capital)),
            ("Change %", round((end["Total Rs"] - capital) / capital * 100, 2)),
            ("Yearly growth % (CAGR)", round(cagr * 100, 2) if cagr is not None else None),
            ("Highest total Rs", f"{high['Total Rs']:,} on {high['Date']}"),
            ("Lowest total Rs", f"{low['Total Rs']:,} on {low['Date']}"),
            ("Lowest trading capital Rs", f"{low_cash['Trading capital Rs']:,} on {low_cash['Date']}"),
            *((f"Trades skipped: {r}", sum(t["Skipped"] == r for t in trades))
              for r in sorted({t["Skipped"] for t in trades} - {""})),
            ("Lots per trade (min / average / max)", lots_used(trades)), ("", "")]  # fmt: skip


def lots_used(trades: list[dict]) -> str:
    n = [t["Lots"] for t in trades if t["Lots"]]
    return f"{min(n)} / {sum(n) / len(n):.1f} / {max(n)}" if n else "none"


def summary(
    trades: list[dict], days: list[dict], capital: float, history: bool
) -> list[tuple]:
    taken = [t for t in trades if t["Lots"]]
    rows: list[tuple] = [("Period", f"{days[0]['Date']} to {days[-1]['Date']}" if days else "no data"),
                         ("Trading days", len(days)), ("Days with a trade", len({t["Date"] for t in taken}))]  # fmt: skip
    for name, picked in (
        ("All", taken),
        *(
            (side, [t for t in taken if t["Side"] == side])
            for side in ("Long", "Short")
        ),
    ):
        rows += [("", ""), (f"{name} trades", len(picked)),
                 (f"{name} win rate %", round(sum(t["Points"] > 0 for t in picked) / len(picked) * 100, 1) if picked else None),
                 (f"{name} {EXITS_LABEL}", " / ".join(str(sum(t["Exit reason"] == r for t in picked)) for r in EXITS)),
                 (f"{name} total points", round(sum(t["Points"] for t in picked), 2)),
                 (f"{name} total R", round(sum(t["R"] for t in picked), 2)),
                 (f"{name} futures Rs", sum(t["Futures Rs"] for t in picked)),
                 (f"{name} option Rs", sum(t["Option Rs"] for t in picked))]  # fmt: skip
    stats = performance([t["Option Rs"] for t in taken], [{"equity": d["Total Rs"] - capital} for d in days], capital,
                        len({t["Date"] for t in taken}))  # fmt: skip
    rows += [("", ""), ("Option trades statistics", "")]
    rows += [
        (k.replace("_", " "), str(v) if isinstance(v, list) else v)
        for k, v in stats.items()
    ]
    data = ("Kaggle NIFTY 1-minute (2015 to Dec 2025) + Groww 5m since, rolled to 5m; VWAP replaced by the session "
            "time-weighted average price (the index has no volume); previous day from the same bars"
            if history else "Yahoo 5m NIFTY (a few minutes delayed), VWAP from NIFTYBEES, previous day from Yahoo daily")  # fmt: skip
    rows += [("", ""), ("Assumptions", ""), ("Data", data),
             ("Filters", f"No trades on expiry days; setups risking under {MIN_RISK} pts skipped"),
             ("Exits", (f"Stop or {R_MULT:g}R target on a 5m high/low (both in one bar = stop), stop to entry from the bar "
                        "after one reaching +1R, else the 15:10 close")),
             ("Option", (f"CE/PE {OTM_STEPS} strikes out of the money, to the expiry in force (monthly Thursday before Feb 2019, weekly Thursday to Aug 2025, "
                        "weekly Tuesday since), lot size of the day, Black-76 at the previous day's India VIX, IV constant")),
             ("Charges", "Today's brokerage and taxes on every trade (older years had different rates)"),
             ("Losing streaks", (f"{LOSS_STREAK} losing trades in a row (index points) = rest of that day and the next "
                                f"{BREAK_DAYS} trading days off. The 2-losses-a-day stop never binds: at most 2 trades a day")),
             ("Not modelled", ("Slippage, bid-ask spread, IV changes, the premium of an overlapping "
                              "second trade (capital is checked one trade at a time)"))]  # fmt: skip
    return rows


# Excel


def write(path: Path, sheets: dict[str, list], capital: float, sizing: str) -> None:
    from openpyxl import (
        Workbook,  # not a project dependency: run with `uv run --with openpyxl`
    )
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for row in sheets.pop("Summary"):
        ws.append(row)
    for cell in ws["A"]:
        cell.font = Font(bold=bool(cell.value) and not str(cell.value)[:1].islower())
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 44, 100
    head, red, green = (
        PatternFill("solid", fgColor="1F4E78"),
        Font(color="C62828"),
        Font(color="2E7D32"),
    )
    for title, rows in sheets.items():
        sheet = wb.create_sheet(title)
        keys = [k for k in (rows[0] if rows else {"Date": 0}) if k != "Why"] + (
            ["Why"] if rows and "Why" in rows[0] else []
        )
        sheet.append(keys)
        for c in sheet[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in rows:
            sheet.append([r[k] for k in keys])
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = (
                60 if col[0].value == "Why" else 14
            )
            if col[0].value in ("Points", "R", "Futures Rs", "Option Rs"):
                for c in col[1:]:
                    c.font = (
                        green
                        if (c.value or 0) > 0
                        else red
                        if (c.value or 0) < 0
                        else Font()
                    )
            if "Rs" in str(col[0].value):
                for c in col[1:]:
                    c.number_format = "#,##0"
        sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
    for title, anchor in (("Daily", "J2"), ("Yearly", "A15")):
        sheet = wb[title]
        cols = [
            i
            for i, c in enumerate(sheet[1], 1)
            if c.value in ("Trading capital Rs", "Wallet Rs", "Total Rs")
        ]
        if sheet.max_row < 2 or not cols:
            continue
        chart = LineChart()
        chart.title = f"Your money from Rs {capital:,.0f}, {sizing}"
        chart.y_axis.title, chart.width, chart.height = "Rs", 26, 12
        for i in cols:
            chart.add_data(
                Reference(sheet, min_col=i, min_row=1, max_row=sheet.max_row),
                titles_from_data=True,
            )
        chart.set_categories(
            Reference(sheet, min_col=1, min_row=2, max_row=sheet.max_row)
        )
        sheet.add_chart(chart, anchor)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(prog="scripts/backtest_liquidity.py")
    parser.add_argument(
        "--history",
        type=Path,
        help="Kaggle NIFTY 1-minute CSV (default: Yahoo's last 60 days)",
    )
    parser.add_argument(
        "--years",
        type=float,
        default=10,
        help="with --history: how many years back from today",
    )
    parser.add_argument("--capital", type=float, default=100_000)
    parser.add_argument("--lots", type=int, default=1)
    parser.add_argument(
        "--book-profits",
        action="store_true",
        help="keep trading capital at --capital and move each win's profit above it to a wallet",
    )
    parser.add_argument(
        "--risk-pct",
        type=float,
        help="percent of the current trading capital at risk per trade: lots = that / option loss per lot at the stop",
    )
    parser.add_argument(
        "--stop",
        choices=("swing", "sweep"),
        default="swing",
        help="stop for the main sheets",
    )
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    settings = get_settings()
    bars, above, daily = history_data(args.history) if args.history else yahoo_data()
    since = (
        (date.today() - timedelta(days=round(365.25 * args.years))).isoformat()
        if args.history
        else ""
    )
    vix_range = f"{int(args.years) + 1}y" if args.history else "1y"
    vix = {
        d["date"][:10]: d["close"]
        for d in parse_yahoo_chart(_chart("^INDIAVIX", "1d", vix_range))
        if d["close"]
    }
    sizing = (
        f"{args.risk_pct:g}% of trading capital at risk per trade"
        if args.risk_pct
        else f"{args.lots} lot{'s' * (args.lots > 1)} per trade"
    )
    runs = {}
    for stop_at in (args.stop, "sweep" if args.stop == "swing" else "swing"):
        trades, days = backtest(settings, bars, above, daily, vix, since, stop_at)
        runs[stop_at] = (
            trades,
            ledger(
                trades,
                days,
                settings,
                args.capital,
                args.lots,
                args.book_profits,
                args.risk_pct,
            ),
        )
    trades, day_rows = runs[args.stop]
    money_rows = money(trades, day_rows, args.capital, sizing, args.book_profits)
    stops = compare(runs, args.capital, sizing, args.book_profits)
    span = f"{args.years:g}y" if args.history else "60d"
    out = (
        args.out
        or Path(__file__).resolve().parents[2]
        / "DocumentsAndDiagrams"
        / f"NIFTY_Liquidity_Backtest_{span}_{date.today()}.xlsx"
    )
    write(out, {"Summary": money_rows + summary(trades, day_rows, args.capital, bool(args.history)),
                "Stops compared": stops, "Yearly": yearly(trades, day_rows), "Trades": trades, "Daily": day_rows},
          args.capital, f"{sizing}, {args.stop} stop")  # fmt: skip
    for row in stops:
        print(" | ".join(str(v) for v in row.values()))
    print("wrote", out)


EXITS, EXITS_LABEL = (
    ("Target", "Stop", "Breakeven", EXIT),
    "targets / stops / breakevens / 15:10",
)
STOP_NAMES = {
    "swing": "Swing stop (current)",
    "sweep": "Stop 5 pts past the sweep high/low",
}


def compare(
    runs: dict[str, tuple[list[dict], list[dict]]],
    capital: float,
    sizing: str,
    book: bool,
) -> list[dict]:
    """One row per metric, one column per stop rule."""
    cols = {}
    for key, (trades, days) in runs.items():
        taken = [t for t in trades if t["Lots"]]
        stats = performance([t["Option Rs"] for t in taken], [{"equity": d["Total Rs"] - capital} for d in days],
                            capital, len({t["Date"] for t in taken}))  # fmt: skip
        cols[STOP_NAMES[key]] = {
            "Signals": len(trades), "Trades taken": len(taken),
            "Win rate %": round(sum(t["Points"] > 0 for t in taken) / len(taken) * 100, 1) if taken else None,
            "Average stop distance (pts)": round(sum(t["Risk pts"] for t in taken) / len(taken), 1) if taken else None,
            EXITS_LABEL.capitalize(): " / ".join(str(sum(t["Exit reason"] == r for t in taken)) for r in EXITS),
            "Total R": round(sum(t["R"] for t in taken), 2),
            "R per trade": round(sum(t["R"] for t in taken) / len(taken), 3) if taken else None,
            "Option Rs": sum(t["Option Rs"] for t in taken),
            "Profit factor": stats["profit_factor"], "Max drawdown %": stats["max_drawdown_pct"],
            "Longest losing streak": stats["max_consecutive_losses"],
            **dict(money(trades, days, capital, sizing, book)[1:-1]),
        }  # fmt: skip
    metrics = dict.fromkeys(k for col in cols.values() for k in col)
    return [
        {"Metric": m, **{name: col.get(m) for name, col in cols.items()}}
        for m in metrics
    ]


if __name__ == "__main__":
    main()
