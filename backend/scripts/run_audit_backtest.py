"""Comprehensive 10-Year Backtest Audit of NIFTY Liquidity Strategies (2015-2026).
Runs and compares:
1. Baseline (PDH/PDL sweeps, 1 lot fixed)
2. Trendline Confluence (PDH/PDL + TL break, 1 lot fixed)
3. Trendline + Equal High/Low Sweeps (EQH/EQL + PDH/PDL, 1 lot fixed)
4. Trendline + Equal + Staged Exits (50% at +2R, +1R BE, 5m candle close trail, 1 lot fixed)
5. Trendline + Equal + Staged Exits + 5% Compounding Risk
6. Strike Selection Comparison: OTM (+2) vs ATM (0) vs ITM (-1)
Outputs full analytics and generates a comprehensive Excel workbook.
"""

from __future__ import annotations

import csv
import json
import sys
import time
from bisect import bisect_left
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.api.market import IST, _chart, parse_yahoo_chart
from app.config import get_settings
from app.engine.metrics import performance
from app.liquidity_signals import OTM_STEPS
from scripts.backtest_liquidity import (
    backtest,
    history_data,
    ledger,
    money,
    yearly,
    write,
)


def calc_stats(trades: list[dict], day_rows: list[dict], capital: float):
    taken = [t for t in trades if t["Lots"]]
    if not taken:
        return {}
    stats = performance(
        [t["Option Rs"] for t in taken],
        [{"equity": d["Total Rs"] - capital} for d in day_rows],
        capital,
        len({t["Date"] for t in taken}),
    )
    longs = [t for t in taken if t["Side"] == "Long"]
    shorts = [t for t in taken if t["Side"] == "Short"]
    wins = [t for t in taken if t["Points"] > 0]
    total_pnl = sum(t["Option Rs"] for t in taken)
    total_pts = sum(t["Points"] for t in taken)
    total_r = sum(t["R"] for t in taken)

    end = day_rows[-1]
    years = (date.fromisoformat(end["Date"]) - date.fromisoformat(day_rows[0]["Date"])).days / 365.25
    cagr = ((end["Total Rs"] / capital) ** (1 / years) - 1) if years > 0 and end["Total Rs"] > 0 else 0

    return {
        "Signals": len(trades),
        "Trades Taken": len(taken),
        "Win Rate %": round(len(wins) / len(taken) * 100, 1),
        "Long Trades (CE)": len(longs),
        "Long Win Rate %": round(sum(t["Points"] > 0 for t in longs) / len(longs) * 100, 1) if longs else 0,
        "Long PnL Rs": sum(t["Option Rs"] for t in longs),
        "Short Trades (PE)": len(shorts),
        "Short Win Rate %": round(sum(t["Points"] > 0 for t in shorts) / len(shorts) * 100, 1) if shorts else 0,
        "Short PnL Rs": sum(t["Option Rs"] for t in shorts),
        "Total Index Points": round(total_pts, 1),
        "Total R": round(total_r, 2),
        "Avg R per Trade": round(total_r / len(taken), 3),
        "Option PnL Rs": total_pnl,
        "Profit Factor": stats.get("profit_factor"),
        "Max Drawdown %": stats.get("max_drawdown_pct"),
        "Max Consecutive Losses": stats.get("max_consecutive_losses"),
        "Ending Capital Rs": end["Total Rs"],
        "Net Profit Rs": end["Total Rs"] - round(capital),
        "Return %": round((end["Total Rs"] - capital) / capital * 100, 1),
        "CAGR %": round(cagr * 100, 2),
    }


def main():
    start_time = time.time()
    settings = get_settings()
    data_path = Path(__file__).resolve().parents[1] / "data" / "intraday" / "NIFTY_1MIN_2015-2025.csv"
    print(f"Loading historical data from {data_path}...")
    bars, above, daily = history_data(data_path)
    print(f"Loaded {len(bars):,} 5m bars across {len(daily):,} days.")

    since = (date.today() - timedelta(days=round(365.25 * 10))).isoformat()
    vix = {
        d["date"][:10]: d["close"]
        for d in parse_yahoo_chart(_chart("^INDIAVIX", "1d", "11y"))
        if d["close"]
    }
    print(f"Loaded {len(vix)} VIX observations. Running backtests from {since}...")

    scenarios = [
        ("Baseline (PDH/PDL)", False, False, False, 2, 100_000, 1, None, False),
        ("Trendline Confluence", True, False, False, 2, 100_000, 1, None, False),
        ("Trendline + EQH/EQL", True, True, False, 2, 100_000, 1, None, False),
        ("Trendline + EQH/EQL + Staged", True, True, True, 2, 100_000, 1, None, False),
        ("Trendline + EQH/EQL + Staged (5% Risk)", True, True, True, 2, 100_000, 1, 5.0, True),
        ("Strike ATM (0 steps)", True, True, True, 0, 100_000, 1, None, False),
        ("Strike ITM (-1 step)", True, True, True, -1, 100_000, 1, None, False),
    ]

    results = {}
    detailed_runs = {}

    for name, req_tl, inc_eq, staged, otm_st, cap, lots, r_pct, book in scenarios:
        t0 = time.time()
        print(f"Running '{name}'...")
        trades, days = backtest(
            settings,
            bars,
            above,
            daily,
            vix,
            since,
            stop_at="swing",
            require_trendline=req_tl,
            include_equal=inc_eq,
            staged=staged,
            otm_steps=otm_st,
            stagnation_bars=18 if staged else 0,
        )
        day_rows = ledger(trades, days, settings, cap, lots, book, risk_pct=r_pct)
        stats = calc_stats(trades, day_rows, cap)
        results[name] = stats
        detailed_runs[name] = (trades, day_rows)
        print(f"  -> Done in {time.time()-t0:.1f}s | Win Rate: {stats.get('Win Rate %')}% | PnL: Rs {stats.get('Option PnL Rs'):,} | PF: {stats.get('Profit Factor')} | DD: {stats.get('Max Drawdown %')}%")

    # Output JSON summary to stdout
    print("\n" + "=" * 80)
    print("ALL 10-YEAR BACKTEST AUDIT RESULTS")
    print("=" * 80)
    print(json.dumps(results, indent=2))

    # Save to Excel
    out_dir = Path(__file__).resolve().parents[2] / "DocumentsAndDiagrams"
    out_path = out_dir / f"NIFTY_Liquidity_Backtest_Full_Audit_10y_{date.today()}.xlsx"
    print(f"\nWriting comprehensive workbook to {out_path}...")

    wb = openpyxl.Workbook()
    ws_exec = wb.active
    ws_exec.title = "Executive Comparison"

    # Header styling
    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    bold_font = Font(bold=True)
    green_font = Font(color="2E7D32")
    red_font = Font(color="C62828")

    # Table of executive comparison
    metrics_list = list(next(iter(results.values())).keys())
    ws_exec.append(["Metric"] + list(results.keys()))
    for c in ws_exec[1]:
        c.fill, c.font = head_fill, head_font

    for m in metrics_list:
        row = [m] + [results[scen].get(m) for scen in results]
        ws_exec.append(row)

    ws_exec.column_dimensions["A"].width = 28
    for col in ws_exec.iter_cols(min_col=2, max_col=len(results) + 1):
        ws_exec.column_dimensions[col[0].column_letter].width = 20
        for cell in col[1:]:
            if isinstance(cell.value, (int, float)):
                if "Rs" in ws_exec.cell(row=cell.row, column=1).value or "Capital" in ws_exec.cell(row=cell.row, column=1).value or "PnL" in ws_exec.cell(row=cell.row, column=1).value or "Profit" in ws_exec.cell(row=cell.row, column=1).value:
                    cell.number_format = "#,##0"
                    cell.font = green_font if cell.value > 0 else red_font if cell.value < 0 else Font()
                elif "Rate" in ws_exec.cell(row=cell.row, column=1).value or "%" in ws_exec.cell(row=cell.row, column=1).value:
                    cell.number_format = "0.0"

    # Add Strike Comparison Sheet
    ws_strike = wb.create_sheet("Strike Selection (OTM-ATM-ITM)")
    strike_scenarios = [
        "Trendline + EQH/EQL + Staged",  # OTM (+2)
        "Strike ATM (0 steps)",
        "Strike ITM (-1 step)",
    ]
    ws_strike.append(["Metric", "OTM (+2 Strikes / ~0.35 Delta)", "ATM (At-The-Money / 0.50 Delta)", "ITM (1 Strike In / ~0.65 Delta)"])
    for c in ws_strike[1]:
        c.fill, c.font = head_fill, head_font
    for m in metrics_list:
        ws_strike.append([m] + [results[s].get(m) for s in strike_scenarios])
    ws_strike.column_dimensions["A"].width = 28
    for col in ws_strike.iter_cols(min_col=2, max_col=4):
        ws_strike.column_dimensions[col[0].column_letter].width = 28
        for cell in col[1:]:
            if isinstance(cell.value, (int, float)):
                if "Rs" in ws_strike.cell(row=cell.row, column=1).value or "PnL" in ws_strike.cell(row=cell.row, column=1).value:
                    cell.number_format = "#,##0"
                    cell.font = green_font if cell.value > 0 else red_font if cell.value < 0 else Font()

    # Add Best Strategy Sheets: Yearly, Trades, Daily
    best_name = "Trendline + EQH/EQL + Staged (5% Risk)"
    best_trades, best_daily = detailed_runs[best_name]

    ws_yearly = wb.create_sheet("Yearly Performance (5% Risk)")
    yr_rows = yearly(best_trades, best_daily)
    if yr_rows:
        ws_yearly.append(list(yr_rows[0].keys()))
        for c in ws_yearly[1]:
            c.fill, c.font = head_fill, head_font
        for r in yr_rows:
            ws_yearly.append(list(r.values()))
        for col in ws_yearly.columns:
            ws_yearly.column_dimensions[col[0].column_letter].width = 16
            if "Rs" in str(col[0].value):
                for cell in col[1:]:
                    cell.number_format = "#,##0"
                    if (cell.value or 0) > 0:
                        cell.font = green_font
                    elif (cell.value or 0) < 0:
                        cell.font = red_font

    ws_trades = wb.create_sheet("All Trades (5% Risk)")
    if best_trades:
        keys = [k for k in best_trades[0].keys() if k != "Why"] + ["Why"]
        ws_trades.append(keys)
        for c in ws_trades[1]:
            c.fill, c.font = head_fill, head_font
        for t in best_trades:
            ws_trades.append([t.get(k) for k in keys])
        for col in ws_trades.columns:
            ws_trades.column_dimensions[col[0].column_letter].width = 50 if col[0].value == "Why" else 14
            if "Rs" in str(col[0].value) or col[0].value in ("Points", "R"):
                for cell in col[1:]:
                    if (cell.value or 0) > 0:
                        cell.font = green_font
                    elif (cell.value or 0) < 0:
                        cell.font = red_font

    ws_daily = wb.create_sheet("Daily Ledger (5% Risk)")
    if best_daily:
        ws_daily.append(list(best_daily[0].keys()))
        for c in ws_daily[1]:
            c.fill, c.font = head_fill, head_font
        for d in best_daily:
            ws_daily.append(list(d.values()))
        for col in ws_daily.columns:
            ws_daily.column_dimensions[col[0].column_letter].width = 18
            if "Rs" in str(col[0].value):
                for cell in col[1:]:
                    cell.number_format = "#,##0"

    wb.save(out_path)
    print(f"Workbook successfully saved to {out_path} in {time.time()-start_time:.1f}s total.")


if __name__ == "__main__":
    main()
