"""uv run python scripts/audit_liquidity.py [--refresh]

Independent audit of the NIFTY PDH/PDL liquidity sweep strategy (app.liquidity_signals). Research only: no orders,
no alerts, nothing in the live path changes.

Data is frozen once in data/intraday/audit_snapshot.json.gz (--refresh rebuilds it): the Kaggle NIFTY 1-minute CSV
rolled to 5m (2015 to 24 Dec 2025) + Groww 5m to END, India VIX (NSE from data/market_eod.db since 2019, Yahoo
before), NSE index EOD bars, and Yahoo's last 60 days of NIFTY and NIFTYBEES 5m bars for the VWAP study. Option
EOD closes are read from data/market_eod.db (NIFTY, 2020 on). Fixed END and seeds: every run gives the same numbers.

The baseline entries are the production scan()'s. Variants come from scan_v(), asserted identical to scan() at its
defaults on every day before any variant is reported. Outputs: DocumentsAndDiagrams/liquidity_strategy_audit/
RESULTS.md (every table), results.json and charts/*.png.
"""

from __future__ import annotations

import argparse
import gzip
import json
import math
import sqlite3
import sys
from bisect import bisect_left, bisect_right
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from statistics import NormalDist

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.api.market import _chart, add_vwap, parse_yahoo_chart
from app.config import get_settings
from app.liquidity_signals import (
    FIRST_ENTRY,
    LAST_ENTRY,
    R_MULT,
    _at,
    _bucket,
    exit_of,
    scan,
    to_15m,
)
from app.market.nse import load_lot_sizes, lot_size_on
from app.paper import order_charges
from app.trading.option_pnl import black76, implied_vol
from scripts.backtest_liquidity import expiry_for, groww_5m, kaggle_5m, twap_above

END, START, IS_END = "2026-10-13", "2016-01-01", "2020-12-31"
SNAP = ROOT / "data" / "intraday" / "audit_snapshot.json.gz"
KAGGLE = ROOT / "data" / "intraday" / "NIFTY_1MIN_2015-2025.csv"
EOD_DB = ROOT / "data" / "market_eod.db"
OUT = ROOT.parent / "DocumentsAndDiagrams" / "liquidity_strategy_audit"
COST_PTS = 4.0  # index pts per round trip: ~Rs 2 of option premium (charges + 1 Rs spread) at delta 0.5
SLIP_RS = 0.5  # option premium Rs per side on top of charges (half of a 1 Rs ATM spread)
SEED = 7


# ---------------------------------------------------------------------------------------------------------- data


def _yahoo(symbol: str, interval: str, range_: str) -> list[dict]:
    keys = ("open", "high", "low", "close")
    return [b for b in parse_yahoo_chart(_chart(symbol, interval, range_)) if None not in (b[k] for k in keys)]


def snapshot(refresh: bool) -> dict:
    if SNAP.exists() and not refresh:
        return json.loads(gzip.decompress(SNAP.read_bytes()))
    bars = kaggle_5m(KAGGLE)
    last = date.fromisoformat(bars[-1]["date"][:10])
    seen = {b["date"] for b in bars}
    bars += [b for b in groww_5m(last + timedelta(days=1), date.fromisoformat(END)) if b["date"] not in seen]
    db = sqlite3.connect(EOD_DB)
    eod = {r[0]: dict(zip(("open", "high", "low", "close", "vix"), r[1:], strict=True)) for r in db.execute(
        "select date, open, high, low, close, vix from index_bars where underlying='NIFTY'")}  # fmt: skip
    snap = {
        "bars": [b for b in bars if b["date"][:10] <= END],
        "groww_overlap": groww_5m(last - timedelta(days=45), last),  # QA: same days from both sources
        "vix_yahoo": {d["date"][:10]: d["close"] for d in _yahoo("^INDIAVIX", "1d", "15y")},
        "eod": eod,
        "y_index": _yahoo("^NSEI", "5m", "60d"),
        "y_bees": _yahoo("NIFTYBEES.NS", "5m", "60d"),
    }
    SNAP.write_bytes(gzip.compress(json.dumps(snap).encode()))
    return snap


def prepare(snap: dict) -> dict:
    """Days, daily bars, VIX and the per-day context (previous day, regime, ATR, expiry), all known before 09:15."""
    settings = get_settings()
    bars = [b for b in snap["bars"] if "09:15" <= b["date"][11:] < "15:30"]
    by_day: dict[str, list[dict]] = defaultdict(list)
    for b in bars:
        by_day[b["date"][:10]].append(b)
    days = sorted(by_day)
    daily = {d: {"open": v[0]["open"], "high": max(b["high"] for b in v), "low": min(b["low"] for b in v),
                 "close": v[-1]["close"], "n": len(v)} for d, v in by_day.items()}  # fmt: skip
    vix = dict(snap["vix_yahoo"]) | {d: r["vix"] for d, r in snap["eod"].items() if r["vix"]}
    vix_dates = sorted(vix)
    trading, closes = set(days), [daily[d]["close"] for d in days]
    tr = [0.0] + [max(daily[d]["high"] - daily[d]["low"], abs(daily[d]["high"] - closes[i - 1]),
                      abs(daily[d]["low"] - closes[i - 1])) for i, d in enumerate(days) if i]  # fmt: skip
    expiry = {d: expiry_for(date.fromisoformat(d), trading, settings).isoformat() for d in days}
    ctx = {}
    for i, d in enumerate(days):
        if i < 70:
            continue
        p = i - 1
        sma = sum(closes[p - 49 : p + 1]) / 50
        sma_then = sum(closes[p - 69 : p - 19]) / 50
        c = closes[p]
        regime = "Bull" if c > sma > sma_then else "Bear" if c < sma < sma_then else "Sideways"
        j = bisect_left(vix_dates, d)
        ctx[d] = {
            "pdh": daily[days[p]]["high"], "pdl": daily[days[p]]["low"], "pdc": c,
            "atr": sum(tr[p - 13 : p + 1]) / 14, "regime": regime, "sma50": sma,
            "vix": vix[vix_dates[j - 1]] if j else None, "expiry": expiry[d],
            "is_expiry": expiry[d] == d, "pre_expiry": i + 1 < len(days) and expiry[d] == days[i + 1],
            "post_expiry": expiry[days[p]] == days[p],
        }  # fmt: skip
    return {"by_day": by_day, "days": days, "daily": daily, "ctx": ctx, "vix": vix, "above": twap_above(bars)}


# ---------------------------------------------------------------------------------------------- research scan

DEFAULT = {"long_tf": "15m", "short_tf": "5m", "long_reclaim": False, "short_reject": True, "disp": 0.0,
           "vwap": True, "open_filter": True, "min_risk": 30.0, "strict": False, "sweep_pts": 1.0,
           "sweep_atr": 0.0, "stop_at": "swing"}  # fmt: skip


def piv(bars: list[dict], key: str, strict: bool) -> list[int]:
    """Same pivots as app.liquidity_signals.pivots over the whole list; `strict` demands a unique extreme. A pivot
    at j uses bars j-2..j+2 only, so the pivots of bars[:m] are exactly these with j <= m - 3."""
    out = []
    for j in range(2, len(bars) - 2):
        w = [bars[x][key] for x in range(j - 2, j + 3)]
        best = max(w) if key == "high" else min(w)
        if bars[j][key] == best and (not strict or w.count(best) == 1):
            out.append(j)
    return out


def _last(pv: list[int], m: int) -> int | None:
    """Last pivot confirmed inside bars[:m]."""
    k = bisect_right(pv, m - 3) - 1
    return pv[k] if k >= 0 else None


def scan_v(bars: list[dict], above: dict, pdh: float, pdl: float, atr: float, cfg: dict | None = None) -> list[dict]:
    """scan() with switchable rules (DEFAULT reproduces it). One rule step per 5m bar; the entry check may run on
    the bar that completes the last step, as in scan()."""
    c = DEFAULT | (cfg or {})
    if not bars:
        return []
    day_open, b15 = bars[0]["open"], to_15m(bars)
    open15: dict[str, float] = {}
    for b in bars:
        open15.setdefault(_bucket(b["date"]), b["open"])
    end15 = {_at(x["date"], 10): k for k, x in enumerate(b15)}
    pv = {(tf, key): piv(bars if tf == "5m" else b15, key, c["strict"]) for tf in ("5m", "15m") for key in ("high", "low")}
    out = []
    for side in ("Long", "Short"):
        long = side == "Long"
        tf = c["long_tf"] if long else c["short_tf"]
        steps = (["reclaim"] if long and c["long_reclaim"] else []) + (["reject"] if not long and c["short_reject"] else []) + ["struct"]
        thr = max(c["sweep_pts"], c["sweep_atr"] * atr)
        stage, ext, swept, struct_at = 0, 0.0, "", ""
        for i, b in enumerate(bars):
            if stage == 0:
                if (long and b["low"] < pdl - thr) or (not long and b["high"] > pdh + thr):
                    stage, ext, swept = 1, b["low"] if long else b["high"], b["date"]
                continue
            ext = min(ext, b["low"]) if long else max(ext, b["high"])
            if stage <= len(steps):
                step, k = steps[stage - 1], end15.get(b["date"])
                if step == "struct":
                    key = "high" if long else "low"
                    if tf == "15m":
                        j = _last(pv[("15m", key)], k) if k is not None else None
                        if j is None:
                            continue
                        level, close, body = b15[j][key], b15[k]["close"], abs(b15[k]["close"] - open15[b15[k]["date"]])
                    else:
                        j = _last(pv[("5m", key)], i)
                        if j is None:
                            continue
                        level, close, body = bars[j][key], b["close"], abs(b["close"] - b["open"])
                    if not (close > level if long else close < level) or body < c["disp"] * atr:
                        continue
                    struct_at = _at(b["date"], 5)
                elif k is None or not (b15[k]["close"] > pdl if step == "reclaim" else b15[k]["close"] < pdh):
                    continue
                stage += 1
                if stage <= len(steps):
                    continue
            closes_at = _at(b["date"], 5)
            if closes_at[11:] > LAST_ENTRY:
                break
            cl = b["close"]
            inside = cl > pdl if long else cl < pdh
            above_open = (cl > day_open if long else cl < day_open) or not c["open_filter"]
            vwap_ok = above.get(b["date"]) == long or not c["vwap"]
            if closes_at[11:] < FIRST_ENTRY or not (inside and above_open and vwap_ok):
                continue
            if c["stop_at"] == "sweep":
                stop = ext - 5 if long else ext + 5
            else:
                key = "low" if long else "high"
                if tf == "15m":
                    done = sum(_at(x["date"], 15) <= closes_at for x in b15)
                    j = _last(pv[("15m", key)], done)
                    sw = b15[j] if j is not None and b15[j]["date"] >= _bucket(swept) else None
                else:
                    j = _last(pv[("5m", key)], i + 1)
                    sw = bars[j] if j is not None and bars[j]["date"] >= swept else None
                if long:
                    stop = sw["low"] - 5 if sw and sw["low"] - 5 < cl else ext - 5
                else:
                    stop = sw["high"] + 5 if sw and sw["high"] + 5 > cl else ext + 15
            risk = abs(cl - stop)
            if risk >= c["min_risk"]:
                out.append({"side": side, "date": b["date"][:10], "time": closes_at, "entry": cl, "stop": round(stop, 2),
                            "risk": round(risk, 2), "level": pdl if long else pdh, "opp": pdh if long else pdl,
                            "ext": ext, "depth": round(pdl - ext if long else ext - pdh, 2), "swept": _at(swept, 5),
                            "struct": struct_at, "day_open": day_open})  # fmt: skip
            break
    return out


# ------------------------------------------------------------------------------------------------------- exits


def run_exit(day: list[dict], s: dict, r_mult: float = 2.0, exit_t: str = "15:10", mode: str = "fixed",
             cost: float = COST_PTS) -> dict:  # fmt: skip
    """Stop first, then target, then the time exit, on 5m highs/lows after the entry bar (as exit_of()). Modes:
    fixed (r_mult target), liquidity (target the opposite previous-day level, at least 1R), be (stop to entry once
    +1R), trail (stop trails 1R behind the best price once +1R), partial (half off at 1R, stop to entry, rest at
    r_mult), structure (no target: exit on a 5m close through the last 5m swing confirmed after entry).
    Stop moves take effect from the next bar: the order of highs and lows inside a bar is unknown."""
    long, e, risk = s["side"] == "Long", s["entry"], s["risk"]
    sg = 1 if long else -1
    stop = s["stop"]
    if mode == "liquidity":
        tgt = s["opp"] if (s["opp"] - e) * sg >= risk else e + sg * risk
    elif mode == "dynamic":
        tgt = e + sg * min(max((s["opp"] - e) * sg, risk), 3 * risk)
    elif mode == "structure":
        tgt = e + sg * 1e9
    else:
        tgt = e + sg * r_mult * risk
    after = [b for b in day if b["date"] >= s["time"]]
    lows_hi = piv(after, "low" if long else "high", False) if mode == "structure" else []
    mfe = mae = 0.0
    part, reason, px, at = None, "Last bar", day[-1]["close"], _at(day[-1]["date"], 5)
    for i, b in enumerate(after):
        at = _at(b["date"], 5)
        fav, adv = (b["high"] - e, e - b["low"]) if long else (e - b["low"], b["high"] - e)
        mae = max(mae, min(adv, risk))
        if (b["low"] <= stop) if long else (b["high"] >= stop):
            reason, px = ("Stop" if (stop - e) * sg < 0 else "Breakeven/trail"), stop
            break
        mfe = max(mfe, min(fav, abs(tgt - e)))
        if (b["high"] >= tgt) if long else (b["low"] <= tgt):
            reason, px = "Target", tgt
            break
        if mode == "structure":
            j = _last(lows_hi, i + 1)
            if j is not None and ((b["close"] < after[j]["low"]) if long else (b["close"] > after[j]["high"])):
                reason, px = "Structure", b["close"]
                break
        if at[11:] >= exit_t:
            reason, px = "Time", b["close"]
            break
        if mode in ("be", "partial") and fav >= risk:
            stop = max(stop, e) if long else min(stop, e)
            if mode == "partial" and part is None:
                part = e + sg * risk
        if mode == "trail" and fav >= risk:
            best = b["high"] if long else b["low"]
            stop = max(stop, best - risk) if long else min(stop, best + risk)
    pts = (px - e) * sg
    if part is not None:
        pts = 0.5 * risk + 0.5 * pts
    t0 = datetime.strptime(s["time"], "%Y-%m-%d %H:%M")
    return {**s, "exit_reason": reason, "exit": round(px, 2), "exit_time": at, "points": round(pts, 2),
            "R_gross": pts / risk, "R": (pts - cost) / risk, "mfe_R": mfe / risk, "mae_R": mae / risk,
            "hold_min": (datetime.strptime(at, "%Y-%m-%d %H:%M") - t0).seconds // 60}  # fmt: skip


# ----------------------------------------------------------------------------------------------------- engine


def cooldown(trades: list[dict], days: list[str], losses: int, off: int) -> list[dict]:
    """After `losses` losing trades in a row (index points, as the live break_until), skip the rest of that day and
    the next `off` trading days."""
    if not losses:
        return trades
    out, streak, paused = [], 0, ""
    for t in sorted(trades, key=lambda t: t["time"]):
        if t["date"] <= paused:
            continue
        out.append(t)
        streak = streak + 1 if t["points"] < 0 else 0
        if streak == losses:
            k = bisect_left(days, t["date"]) + off
            streak, paused = 0, days[k] if k < len(days) else "9999"
    return out


def run(D: dict, cfg: dict | None = None, exit_kw: dict | None = None, expiry: str = "skip", cool=(3, 3),
        day_filter=None, since: str = START, until: str = END) -> list[dict]:  # fmt: skip
    """Every trade of a configuration. expiry: skip | only | keep. day_filter(ctx, signal) -> keep?"""
    trades = []
    for d in D["days"]:
        x = D["ctx"].get(d)
        if x is None or not since <= d <= until:
            continue
        if (expiry == "skip" and x["is_expiry"]) or (expiry == "only" and not x["is_expiry"]):
            continue
        for s in scan_v(D["by_day"][d], D["above"], x["pdh"], x["pdl"], x["atr"], cfg):
            if day_filter is None or day_filter(x, s):
                trades.append(run_exit(D["by_day"][d], s, **(exit_kw or {})) | {"ctx": x})
    days = [d for d in D["days"] if since <= d <= until]
    return cooldown(trades, days, *cool) if cool else sorted(trades, key=lambda t: t["time"])


def bootstrap_ci(x: list[float], n: int = 4000) -> tuple[float, float]:
    if len(x) < 2:
        return (float("nan"), float("nan"))
    a = np.asarray(x)
    means = np.random.default_rng(SEED).choice(a, (n, len(a))).mean(axis=1)
    return float(np.percentile(means, 2.5)), float(np.percentile(means, 97.5))


def stats(trades: list[dict], key: str = "R", days: int | None = None) -> dict:
    r = [t[key] for t in trades]
    n = len(r)
    if not n:
        return {"Trades": 0}
    wins, losses = [x for x in r if x > 0], [x for x in r if x <= 0]
    eq = np.cumsum(r)
    dd = float(np.max(np.maximum.accumulate(np.concatenate([[0], eq]))[1:] - eq)) if n else 0.0
    streak = longest = wstreak = wlongest = 0
    for x in r:
        streak = streak + 1 if x <= 0 else 0
        wstreak = wstreak + 1 if x > 0 else 0
        longest, wlongest = max(longest, streak), max(wlongest, wstreak)
    sd = float(np.std(r, ddof=1)) if n > 1 else 0.0
    lo, hi = bootstrap_ci(r)
    years = max((date.fromisoformat(trades[-1]["date"]) - date.fromisoformat(trades[0]["date"])).days / 365.25, 0.5)
    by_day = defaultdict(float)
    for t in trades:
        by_day[t["date"]] += t[key]
    ndays = days or round(years * 248)
    daily = np.array(list(by_day.values()) + [0.0] * max(ndays - len(by_day), 0))
    dsd, dn = daily.std(ddof=1), math.sqrt(np.mean(np.minimum(daily, 0) ** 2))
    return {
        "Trades": n, "Trades/yr": round(n / years, 1), "Win %": round(len(wins) / n * 100, 1),
        "Avg win R": round(float(np.mean(wins)), 3) if wins else None,
        "Avg loss R": round(float(np.mean(losses)), 3) if losses else None,
        "Gross profit R": round(sum(wins), 2), "Gross loss R": round(sum(losses), 2),
        "Expectancy R": round(float(np.mean(r)), 4), "CI95 low": round(lo, 4), "CI95 high": round(hi, 4),
        "t-stat": round(float(np.mean(r)) / (sd / math.sqrt(n)), 2) if sd else None,
        "Total R": round(float(eq[-1]), 2), "PF": round(sum(wins) / -sum(losses), 3) if losses and sum(losses) else None,
        "Max DD R": round(dd, 2), "Recovery factor": round(float(eq[-1]) / dd, 2) if dd else None,
        "Max consec losses": longest, "Max consec wins": wlongest,
        "Sharpe (daily, ann.)": round(float(daily.mean() / dsd * math.sqrt(248)), 2) if dsd else None,
        "Sortino (ann.)": round(float(daily.mean() / dn * math.sqrt(248)), 2) if dn else None,
        "Avg hold min": round(float(np.mean([t["hold_min"] for t in trades])), 0),
        "Avg MFE R": round(float(np.mean([t["mfe_R"] for t in trades])), 2),
        "Avg MAE R": round(float(np.mean([t["mae_R"] for t in trades])), 2),
        "Reached 1R %": round(sum(t["mfe_R"] >= 1 for t in trades) / n * 100, 1),
        "Reached 2R %": round(sum(t["mfe_R"] >= 2 - 1e-9 for t in trades) / n * 100, 1),
        "Stopped %": round(sum(t["exit_reason"] == "Stop" for t in trades) / n * 100, 1),
        "1R then stopped %": round(sum(t["mfe_R"] >= 1 and t["exit_reason"] in ("Stop", "Breakeven/trail") for t in trades) / n * 100, 1),
    }  # fmt: skip


SHORT_COLS = ["Trades", "Trades/yr", "Win %", "Expectancy R", "CI95 low", "CI95 high", "t-stat", "Total R", "PF",
              "Max DD R", "Max consec losses"]  # fmt: skip


def split(trades: list[dict]) -> dict:
    return {"IS 2016-2020": [t for t in trades if t["date"] <= IS_END], "OOS 2021-2026": [t for t in trades if t["date"] > IS_END]}


def md(rows: list[dict], cols: list[str] | None = None) -> str:
    if not rows:
        return "_no rows_\n"
    cols = cols or list(rows[0])
    fmt = lambda v: f"{v:,.4g}" if isinstance(v, float) else ("" if v is None else str(v))
    return "\n".join(["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols),
                      *("| " + " | ".join(fmt(r.get(c)) for c in cols) + " |" for r in rows)]) + "\n"  # fmt: skip


def row(name: str, trades: list[dict], cols=SHORT_COLS, **extra) -> dict:
    s = stats(trades)
    return {"Config": name, **extra, **{c: s.get(c) for c in cols}}


# ------------------------------------------------------------------------------------------------ option model

LOTS = load_lot_sizes(ROOT / "data" / "lot_sizes" / "nse.yaml")


def lot_on(d: str) -> int:
    return lot_size_on(LOTS, "NIFTY", d) or 75  # 75 from Nov 2015 until the table starts in 2019


def _years(expiry: str, at: str) -> float:
    left = datetime.fromisoformat(f"{expiry} 15:30") - datetime.strptime(at, "%Y-%m-%d %H:%M")
    return max(left.total_seconds(), 0) / 86400 / 365


def option_trade(t: dict, settings, iv: float | None = None, basis: float = 0.0, slip: float = SLIP_RS,
                 iv_exit: float = 1.0, expiry: str | None = None, otm: int = 0) -> dict:  # fmt: skip
    """1 lot of the CE (long) / PE (short) `otm` strikes out of the money (0 = ATM) bought at the entry bar close and
    sold at the index exit: Black-76, forward = index + basis, IV `iv` (default previous-day India VIX) held constant
    (x iv_exit at the exit), charges of the trade date and `slip` Rs worse per side. R = net P&L / planned net loss
    at the index stop."""
    long = t["side"] == "Long"
    kind, k = ("CE" if long else "PE"), (round(t["entry"] / 50) + (otm if long else -otm)) * 50
    exp = expiry or t["ctx"]["expiry"]
    vol = iv if iv is not None else t["ctx"]["vix"] / 100
    lot = lot_on(t["date"])
    p_in = black76(t["entry"] + basis, k, _years(exp, t["time"]), vol, kind) + slip
    p_out = max(black76(t["exit"] + basis, k, _years(exp, t["exit_time"]), vol * iv_exit, kind) - slip, 0.05)
    p_stop = max(black76(t["stop"] + basis, k, max(_years(exp, t["time"]) - 0.15 / 365, 0), vol, kind) - slip, 0.05)
    ch = lambda side, p: order_charges(settings, kind, side, p, lot, t["date"])["total"]
    pnl = (p_out - p_in) * lot - ch("BUY", p_in) - ch("SELL", p_out)
    loss = (p_in - p_stop) * lot + ch("BUY", p_in) + ch("SELL", p_stop)
    return {"premium_in": p_in, "premium_out": p_out, "premium_stop": p_stop, "lot": lot, "pnl": pnl,
            "loss_at_stop": loss, "opt_R": pnl / loss if loss > 0 else float("nan")}  # fmt: skip


class Eod:
    """NSE NIFTY option end-of-day closes (2020 on) from data/market_eod.db."""

    def __init__(self):
        self.db = sqlite3.connect(EOD_DB)
        self.expiries: dict[str, list[str]] = defaultdict(list)
        for d, e in self.db.execute("select distinct date, expiry from option_eod where underlying='NIFTY' order by date, expiry"):
            self.expiries[d].append(e)

    def nearest(self, d: str) -> str | None:
        return next((e for e in self.expiries.get(d, []) if e >= d), None)

    def quotes(self, d: str, e: str) -> dict[tuple[float, str], float]:
        q = self.db.execute("select strike, option_type, close from option_eod where underlying='NIFTY' and date=? "
                            "and expiry=? and close > 0", (d, e))  # fmt: skip
        return {(k, t): c for k, t, c in q}

    def contract_iv(self, d: str, e: str, k: float, kind: str, spot: float) -> tuple[float, float] | None:
        """(IV, forward basis) of strike `k` at the close of `d`, forward from put-call parity at that strike."""
        q = self.quotes(d, e)
        if (k, "CE") not in q or (k, "PE") not in q:
            return None
        f = k + q[(k, "CE")] - q[(k, "PE")]
        iv = implied_vol(q[(k, kind)], f, k, _years(e, f"{d} 15:30"), kind)
        return (iv, f - spot) if iv else None


# ------------------------------------------------------------------------------------------------- analyses


def data_quality(D: dict, snap: dict) -> dict:
    days, daily = D["days"], D["daily"]
    years = Counter(d[:4] for d in days)
    short = [(d, daily[d]["n"]) for d in days if daily[d]["n"] < 75]
    weekend = [d for d in days if date.fromisoformat(d).weekday() >= 5]
    diffs = defaultdict(list)
    for d in days:
        e = snap["eod"].get(d)
        if e and e["high"] and d <= "2025-12-24":
            for k in ("high", "low", "close"):
                diffs[k].append(daily[d][k] - e[k])
    eod_cmp = [{"Field": k, "Days": len(v), "Mean diff": round(float(np.mean(v)), 2),
                "Median abs diff": round(float(np.median(np.abs(v))), 2), "P99 abs diff": round(float(np.percentile(np.abs(v), 99)), 2),
                "Days with abs diff > 5 pts": int(sum(abs(x) > 5 for x in v))} for k, v in diffs.items()]  # fmt: skip
    kag = {b["date"]: b for b in snap["bars"] if b["date"][:10] <= "2025-12-24"}
    ov = [(kag[b["date"]], b) for b in snap["groww_overlap"] if b["date"] in kag]
    gv = [{"Field": k, "Bars": len(ov), "Median abs diff": round(float(np.median([abs(a[k] - b[k]) for a, b in ov])), 2),
           "Max abs diff": round(max(abs(a[k] - b[k]) for a, b in ov), 2)} for k in ("open", "high", "low", "close")] if ov else []  # fmt: skip
    missing_eod = sorted(d for d in snap["eod"] if START <= d <= "2025-12-24" and d not in daily)
    extra = sorted(d for d in days if "2019-01-01" <= d <= "2025-12-24" and d not in snap["eod"])
    return {"years": [{"Year": y, "Sessions": n} for y, n in sorted(years.items())], "short": short, "weekend": weekend,
            "eod_cmp": eod_cmp, "groww_cmp": gv, "missing_vs_eod": missing_eod, "extra_vs_eod": extra,
            "first": days[0], "last": days[-1], "bars": sum(daily[d]["n"] for d in days)}  # fmt: skip


def spec_exit_of(day: list[dict], s: dict) -> tuple[str, float, str]:
    """exit_of() as specified (and as in the working tree when the audit began): 2R, no stop move, 15:10."""
    long = s["side"] == "Long"
    target = s["entry"] + 2 * s["risk"] if long else s["entry"] - 2 * s["risk"]
    for b in day:
        if b["date"] < s["time"]:
            continue
        at = _at(b["date"], 5)
        if b["low"] <= s["stop"] if long else b["high"] >= s["stop"]:
            return "Stop", s["stop"], at
        if b["high"] >= target if long else b["low"] <= target:
            return "Target", target, at
        if at[11:] >= "15:10":
            return "15:10", b["close"], at
    return "Last bar", day[-1]["close"], _at(day[-1]["date"], 5)


def reproduce(D: dict) -> dict:
    """Production scan() on every day vs scan_v defaults; run_exit vs the specified exit and vs the working tree's
    exit_of() (R_MULT target, stop to entry at +1R); and a no-repaint check: each signal again from the bars up to
    its entry bar only (and absent one bar earlier)."""
    mism, n, exit_mism, wt_mism, repaint, early = [], 0, 0, 0, 0, 0
    for d in D["days"]:
        x = D["ctx"].get(d)
        if x is None or d < START or x["is_expiry"]:
            continue
        day = D["by_day"][d]
        prod = scan(day, D["above"], x["pdh"], x["pdl"])
        mine = scan_v(day, D["above"], x["pdh"], x["pdl"], x["atr"])
        key = lambda s: (s["side"], s["time"], s["entry"], s["stop"], s["risk"])
        if [key(s) for s in prod] != [key(s) for s in mine]:
            mism.append(d)
        for s in prod:
            n += 1
            for ref, kw in ((spec_exit_of, {}), (exit_of, {"r_mult": R_MULT, "mode": "be"})):
                _reason, px, at = ref(day, s)
                r = run_exit(day, {**s, "date": d, "opp": 0}, cost=0, **kw)
                if (round(px, 2), at) != (round(r["exit"], 2), r["exit_time"]):
                    exit_mism += ref is spec_exit_of
                    wt_mism += ref is exit_of
            cut = [b for b in day if _at(b["date"], 5) <= s["time"]]
            again = [key(z) for z in scan(cut, D["above"], x["pdh"], x["pdl"]) if z["side"] == s["side"]]
            repaint += again != [key(s)]
            early += any(z["side"] == s["side"] for z in scan(cut[:-1], D["above"], x["pdh"], x["pdl"]))
    return {"signals": n, "scan_mismatch_days": mism, "exit_mismatches": exit_mism, "wt_exit_mismatches": wt_mism,
            "repaint": repaint, "early": early}  # fmt: skip


def option_validation(D: dict, eod: Eod) -> dict:
    """Black-76 at India VIX (forward = spot, the live fallback) vs the real NSE close of the nearest-expiry ATM
    option, 2020 on, at the 15:30 close. Also the contract's own IV vs VIX, the parity forward basis and skew."""
    rows = []
    for d in sorted(eod.expiries):
        e, spot = eod.nearest(d), D["daily"].get(d, {}).get("close")
        x = D["ctx"].get(d)
        if not e or e == d or not spot or not x or not x["vix"]:
            continue
        q, yrs = eod.quotes(d, e), _years(e, f"{d} 15:30")
        k = round(spot / 50) * 50
        if (k, "CE") not in q or (k, "PE") not in q:
            continue
        f = k + q[(k, "CE")] - q[(k, "PE")]
        vix_today = D["vix"].get(d)
        dte = (date.fromisoformat(e) - date.fromisoformat(d)).days
        r = {"date": d, "dte": dte, "basis": f - spot}
        for kind in ("CE", "PE"):
            actual = q[(k, kind)]
            iv = implied_vol(actual, f, k, yrs, kind)
            r[kind] = {"actual": actual, "model_prev": black76(spot, k, yrs, x["vix"] / 100, kind),
                       "model_today": black76(spot, k, yrs, vix_today / 100, kind) if vix_today else None,
                       "model_fwd": black76(f, k, yrs, vix_today / 100, kind) if vix_today else None,
                       "iv": iv, "vix": vix_today}  # fmt: skip
        lo = max((s for s, t in q if t == "PE" and s <= spot * 0.98), default=None)
        hi = min((s for s, t in q if t == "CE" and s >= spot * 1.02), default=None)
        if lo and hi and r["PE"]["iv"]:
            ivl = implied_vol(q[(lo, "PE")], f, lo, yrs, "PE")
            ivh = implied_vol(q[(hi, "CE")], f, hi, yrs, "CE")
            atm = (r["PE"]["iv"] + (r["CE"]["iv"] or r["PE"]["iv"])) / 2
            r["skew"] = (ivl - atm, ivh - atm) if ivl and ivh else None
        rows.append(r)
    buckets = [("1", 1, 1), ("2-3", 2, 3), ("4-7", 4, 7), ("8-35", 8, 35)]
    table = []
    for kind in ("CE", "PE"):
        for name, lo_, hi_ in buckets:
            sel = [r[kind] for r in rows if lo_ <= r["dte"] <= hi_ and r[kind]["model_today"]]
            if not sel:
                continue
            err = [(s["model_prev"] - s["actual"]) / s["actual"] * 100 for s in sel]
            err_t = [(s["model_today"] - s["actual"]) / s["actual"] * 100 for s in sel]
            err_f = [(s["model_fwd"] - s["actual"]) / s["actual"] * 100 for s in sel]
            ratio = [s["iv"] / (s["vix"] / 100) for s in sel if s["iv"]]
            table.append({"Type": kind, "Days to expiry": name, "Days": len(sel),
                          "Median err % (prev-day VIX, F=spot)": round(float(np.median(err)), 1),
                          "Mean abs err % (prev-day VIX)": round(float(np.mean(np.abs(err))), 1),
                          "P90 abs err % (prev-day VIX)": round(float(np.percentile(np.abs(err), 90)), 1),
                          "Median err % (same-day VIX, F=spot)": round(float(np.median(err_t)), 1),
                          "Median err % (same-day VIX, parity F)": round(float(np.median(err_f)), 1),
                          "Median contract IV / VIX": round(float(np.median(ratio)), 2) if ratio else None})  # fmt: skip
    basis = [r["basis"] for r in rows]
    skew = [r["skew"] for r in rows if r.get("skew")]
    return {"table": table, "days": len(rows), "basis_median": float(np.median(basis)), "basis_p10": float(np.percentile(basis, 10)),
            "basis_p90": float(np.percentile(basis, 90)),
            "skew_put_2pct": float(np.median([s[0] for s in skew])) if skew else None,
            "skew_call_2pct": float(np.median([s[1] for s in skew])) if skew else None, "rows": rows}  # fmt: skip


def expiry_check(D: dict, eod: Eod) -> list[dict]:
    """expiry_for() (used by the backtest and its expiry filter) vs the nearest expiry NSE actually listed."""
    out = []
    for d in D["days"]:
        x, real = D["ctx"].get(d), eod.nearest(d)
        if x and real and x["expiry"] != real:
            out.append({"Date": d, "expiry_for()": x["expiry"], "NSE nearest listed": real})
    return out


def vwap_study(snap: dict) -> dict:
    """Last 60 days: NIFTYBEES close vs its VWAP (live) against the index vs its TWAP (10-year backtest) and the
    ETF vs its own TWAP; agreement on every bar and on the bars near the line; signals under each source."""
    idx = {b["date"]: b for b in snap["y_index"] if "09:15" <= b["date"][11:] < "15:30"}
    bees = add_vwap([dict(b) for b in snap["y_bees"]])
    live = {b["date"]: b["close"] > b["vwap"] for b in bees if b["vwap"]}
    bees_twap = twap_above([b for b in bees if "09:15" <= b["date"][11:] < "15:30"])
    idx_bars = sorted(idx.values(), key=lambda b: b["date"])
    twap = twap_above(idx_bars)
    common = [d for d in twap if d in live]
    mean_tp, running = {}, defaultdict(lambda: [0.0, 0])
    for b in idx_bars:
        acc = running[b["date"][:10]]
        acc[0] += (b["high"] + b["low"] + b["close"]) / 3
        acc[1] += 1
        mean_tp[b["date"]] = acc[0] / acc[1]
    near = [d for d in common if abs(idx[d]["close"] - mean_tp[d]) / idx[d]["close"] < 0.001]
    agree = lambda a, b, ds: round(sum(a[d] == b.get(d) for d in ds) / len(ds) * 100, 1) if ds else None
    by_day = defaultdict(list)
    for b in idx_bars:
        by_day[b["date"][:10]].append(b)
    ds = sorted(by_day)
    sig = {"NIFTYBEES VWAP (live)": [], "Index TWAP (10y backtest)": []}
    for i, d in enumerate(ds[1:], 1):
        p = by_day[ds[i - 1]]
        pdh, pdl = max(b["high"] for b in p), min(b["low"] for b in p)
        for name, src in (("NIFTYBEES VWAP (live)", live), ("Index TWAP (10y backtest)", twap)):
            sig[name] += [(d, s["side"], s["time"][11:], s["entry"]) for s in scan(by_day[d], src, pdh, pdl)]
    missing = sum(d not in live for d in idx)
    return {"bars": len(common), "days": len(ds), "agree_all": agree(twap, live, common), "near_bars": len(near),
            "agree_near": agree(twap, live, near), "agree_bees_twap": agree(bees_twap, live, [d for d in common if d in bees_twap]),
            "index_bars_without_bees_vwap": missing, "signals": sig,
            "only_live": sorted(set(sig["NIFTYBEES VWAP (live)"]) - set(sig["Index TWAP (10y backtest)"])),
            "only_twap": sorted(set(sig["Index TWAP (10y backtest)"]) - set(sig["NIFTYBEES VWAP (live)"]))}  # fmt: skip


def sweep_outcomes(D: dict) -> list[dict]:
    """Every PDL/PDH sweep (not only trades): does the day close back inside the level (reversal) or beyond it
    (continuation)? Depth is the day's full excursion beyond the level: descriptive, not tradable."""
    buckets = [("1-2", 1, 2), ("2-5", 2, 5), ("5-10", 5, 10), ("10-20", 10, 20), (">20", 20, 1e9)]
    out = []
    for side in ("PDL sweep (long setup)", "PDH raid (short setup)"):
        long = side.startswith("PDL")
        for name, lo, hi in buckets:
            n = back = 0
            for d in D["days"]:
                x = D["ctx"].get(d)
                if not x or d < START:
                    continue
                v = D["daily"][d]
                depth = x["pdl"] - v["low"] if long else v["high"] - x["pdh"]
                if lo < depth <= hi:
                    n += 1
                    back += (v["close"] > x["pdl"]) if long else (v["close"] < x["pdh"])
            out.append({"Setup": side, "Depth pts": name, "Days": n, "Closed back inside %": round(back / n * 100, 1) if n else None})
    return out


def sizing(trades: list[dict], capital: float, pct: float) -> dict:
    """Compounding, whole lots: lots = pct% of equity / modelled option loss per lot at the stop (charges in);
    0 lots = skipped; never more premium than the equity."""
    eq, peak, max_dd, skipped, taken, worst_day, outlay = capital, capital, 0.0, 0, 0, 0.0, 0.0
    curve, day_pnl, dd_start, longest = [], defaultdict(float), None, 0
    for t in trades:
        o = t["opt"]
        lots = int(eq * pct / 100 // o["loss_at_stop"]) if o["loss_at_stop"] > 0 else 0
        lots = min(lots, int(eq // (o["premium_in"] * o["lot"])))
        if lots < 1:
            skipped += 1
            continue
        taken += 1
        outlay = max(outlay, lots * o["premium_in"] * o["lot"] / eq * 100)
        day_pnl[t["date"]] += lots * o["pnl"] / eq * 100
        eq += lots * o["pnl"]
        if eq > peak:
            if dd_start:
                longest = max(longest, (date.fromisoformat(t["date"]) - date.fromisoformat(dd_start)).days)
            peak, dd_start = eq, None
        elif dd_start is None:
            dd_start = t["date"]
        max_dd = max(max_dd, (peak - eq) / peak * 100)
        curve.append((t["date"], eq))
    if dd_start:
        longest = max(longest, (date.fromisoformat(trades[-1]["date"]) - date.fromisoformat(dd_start)).days)
    yrs = (date.fromisoformat(trades[-1]["date"]) - date.fromisoformat(trades[0]["date"])).days / 365.25
    worst_day = min(day_pnl.values(), default=0.0)
    return {"Capital Rs": f"{capital:,.0f}", "Risk %": pct, "Taken": taken, "Skipped (0 lots)": skipped,
            "End Rs": round(eq), "CAGR %": round(((eq / capital) ** (1 / yrs) - 1) * 100, 1) if eq > 0 else -100,
            "Max DD %": round(max_dd, 1), "Longest DD days": longest, "Worst day %": round(worst_day, 1),
            "Max premium outlay % of equity": round(outlay, 1), "curve": curve}  # fmt: skip


def ruin(r: np.ndarray, pct: float, n: int, sims: int = 4000) -> dict:
    """Bootstrap `n`-trade paths of per-trade R (fractional sizing, no lot rounding): P(drawdown >= 30%/50%)."""
    g = np.random.default_rng(SEED)
    paths = np.cumprod(1 + g.choice(r, (sims, n)) * pct / 100, axis=1)
    paths = np.concatenate([np.ones((sims, 1)), paths], axis=1)
    dd = 1 - (paths / np.maximum.accumulate(paths, axis=1)).min(axis=1)
    return {"Risk %": pct, "P(DD>=20%)": round(float((dd >= 0.2).mean() * 100), 1), "P(DD>=30%)": round(float((dd >= 0.3).mean() * 100), 1),
            "P(DD>=50%)": round(float((dd >= 0.5).mean() * 100), 2), "Median DD %": round(float(np.median(dd) * 100), 1),
            "P95 DD %": round(float(np.percentile(dd, 95) * 100), 1), "P(end below start)": round(float((paths[:, -1] < 1).mean() * 100), 1)}  # fmt: skip


def walk_forward(D: dict, grid: dict[str, list], make, base: list[dict]) -> dict:
    """Anchored yearly walk-forward: each year from 2019 uses the grid value with the best expectancy on all
    earlier years (>= 30 trades), then trades that year with it."""
    runs = {v: make(v) for v in grid["values"]}
    years = [str(y) for y in range(2019, 2027)]
    picked, oos = [], []
    for y in years:
        best, best_e = None, -1e9
        for v, tr in runs.items():
            past = [t["R"] for t in tr if t["date"][:4] < y]
            if len(past) >= 30 and np.mean(past) > best_e:
                best, best_e = v, float(np.mean(past))
        picked.append({"Year": y, grid["name"]: best, "IS expectancy R": round(best_e, 3)})
        oos += [t for t in runs[best] if t["date"][:4] == y]
    base_y = [t for t in base if t["date"][:4] >= "2019"]
    return {"picked": picked, "wf": stats(oos), "base": stats(base_y)}


# ----------------------------------------------------------------------------------------------------- charts


def charts(res: dict, base: list[dict], opt_rows: list[dict], sizes: list[dict], mc_dd: np.ndarray) -> list[str]:
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    out = OUT / "charts"
    out.mkdir(parents=True, exist_ok=True)
    names = []

    def save(fig, name):
        fig.tight_layout()
        fig.savefig(out / name, dpi=110)
        plt.close(fig)
        names.append(name)

    dates = [datetime.fromisoformat(t["date"]) for t in base]
    eq = np.cumsum([t["R"] for t in base])
    eq_g = np.cumsum([t["R_gross"] for t in base])
    fig, ax = plt.subplots(figsize=(11, 4.5))
    ax.plot(dates, eq_g, label="Gross (no costs)", color="#90a4ae")
    ax.plot(dates, eq, label=f"Net of {COST_PTS:g} index pts/trade", color="#1e88e5")
    ax.axvline(datetime.fromisoformat("2021-01-01"), color="k", ls=":", lw=1)
    ax.text(datetime.fromisoformat("2021-02-01"), ax.get_ylim()[1] * 0.9, "out-of-sample ->", fontsize=8)
    ax.set_title("Original strategy: cumulative R (index), 2016-2026")
    ax.set_ylabel("R")
    ax.legend()
    ax.grid(alpha=0.3)
    save(fig, "01_equity_curve.png")

    dd = np.maximum.accumulate(np.concatenate([[0], eq]))[1:] - eq
    fig, ax = plt.subplots(figsize=(11, 3.5))
    ax.fill_between(dates, -dd, color="#e53935", alpha=0.5)
    ax.set_title("Original strategy: drawdown in R (net)")
    ax.set_ylabel("R")
    ax.grid(alpha=0.3)
    save(fig, "02_drawdown.png")

    months = defaultdict(float)
    for t in base:
        months[t["date"][:7]] += t["R"]
    ys = sorted({m[:4] for m in months})
    grid = np.array([[months.get(f"{y}-{m:02d}", 0.0) for m in range(1, 13)] for y in ys])
    fig, ax = plt.subplots(figsize=(10, 5))
    lim = np.abs(grid).max() or 1
    im = ax.imshow(grid, cmap="RdYlGn", vmin=-lim, vmax=lim, aspect="auto")
    ax.set_xticks(range(12), ["J", "F", "M", "A", "M", "J", "J", "A", "S", "O", "N", "D"])
    ax.set_yticks(range(len(ys)), ys)
    for i in range(len(ys)):
        for j in range(12):
            if grid[i, j]:
                ax.text(j, i, f"{grid[i, j]:.1f}", ha="center", va="center", fontsize=7)
    fig.colorbar(im, label="R")
    ax.set_title("Monthly net R (1R = 1% of capital at 1% risk)")
    save(fig, "03_monthly_returns.png")

    yr = res["by_year"]
    fig, ax = plt.subplots(figsize=(10, 4))
    ax.bar([r["Year"] for r in yr], [r["Total R"] for r in yr], color=["#43a047" if r["Total R"] > 0 else "#e53935" for r in yr])
    ax.set_title("Annual net R")
    ax.grid(alpha=0.3, axis="y")
    save(fig, "04_annual_returns.png")

    fig, ax = plt.subplots(figsize=(11, 4.5))
    for side, col in (("Long", "#43a047"), ("Short", "#e53935")):
        tr = [t for t in base if t["side"] == side]
        ax.plot([datetime.fromisoformat(t["date"]) for t in tr], np.cumsum([t["R"] for t in tr]), color=col,
                label=f"{side} ({'PDL sweep' if side == 'Long' else 'PDH raid'}), n={len(tr)}")  # fmt: skip
    ax.set_title("Long (PDL sweep) vs short (PDH raid): cumulative net R")
    ax.legend()
    ax.grid(alpha=0.3)
    save(fig, "05_long_vs_short_pdl_vs_pdh.png")

    for fname, key, title in (("06_by_entry_time.png", "by_time", "Expectancy by entry time"),
                              ("07_by_volatility.png", "by_vix", "Expectancy by India VIX regime (previous day)")):  # fmt: skip
        rows = res[key]
        fig, ax = plt.subplots(figsize=(9, 4))
        x = range(len(rows))
        ax.bar(x, [r["Expectancy R"] or 0 for r in rows], color="#5c6bc0")
        ax.errorbar(x, [r["Expectancy R"] or 0 for r in rows],
                    yerr=[[max((r["Expectancy R"] or 0) - (r["CI95 low"] or 0), 0) for r in rows],
                          [max((r["CI95 high"] or 0) - (r["Expectancy R"] or 0), 0) for r in rows]], fmt="none", color="k", capsize=4)  # fmt: skip
        ax.set_xticks(list(x), [f"{r['Bucket']}\n(n={r['Trades']})" for r in rows], fontsize=8)
        ax.axhline(0, color="k", lw=0.8)
        ax.set_title(title + " (net R, 95% bootstrap CI)")
        ax.grid(alpha=0.3, axis="y")
        save(fig, fname)

    fig, axs = plt.subplots(1, 2, figsize=(11, 4))
    for ax, side in zip(axs, ("Long", "Short"), strict=True):
        ax.hist([t["depth"] for t in base if t["side"] == side], bins=40, color="#7e57c2")
        ax.set_title(f"{side}: sweep depth at entry (pts beyond level)")
    save(fig, "08_sweep_depth.png")

    fig, axs = plt.subplots(1, 2, figsize=(11, 4))
    axs[0].hist([t["risk"] for t in base], bins=40, color="#ef6c00")
    axs[0].set_title("Initial risk (index pts)")
    reasons = Counter(t["exit_reason"] for t in base)
    axs[1].bar(list(reasons), list(reasons.values()), color="#00897b")
    axs[1].set_title("Exit reason")
    save(fig, "09_stop_target_distribution.png")

    fig, ax = plt.subplots(figsize=(7, 6))
    col = ["#43a047" if t["R"] > 0 else "#e53935" for t in base]
    ax.scatter([t["mae_R"] for t in base], [t["mfe_R"] for t in base], c=col, s=10, alpha=0.6)
    ax.set_xlabel("MAE (R)")
    ax.set_ylabel("MFE (R)")
    ax.set_title("MAE vs MFE (green = net winner)")
    ax.grid(alpha=0.3)
    save(fig, "10_mae_mfe.png")

    fig, ax = plt.subplots(figsize=(7, 6))
    ax.scatter([t["R_gross"] for t in opt_rows], [t["opt"]["opt_R"] for t in opt_rows], s=10, alpha=0.5, color="#1e88e5")
    ax.plot([-1.5, 2.5], [-1.5, 2.5], color="k", lw=0.8, ls=":")
    ax.set_xlabel("Index R (gross)")
    ax.set_ylabel("Option R (modelled, charges + slippage)")
    ax.set_title("Option vs underlying return per trade")
    ax.grid(alpha=0.3)
    save(fig, "11_option_vs_underlying.png")

    fig, ax = plt.subplots(figsize=(11, 4.5))
    for s in sizes:
        if s["Capital Rs"] == "500,000":
            ax.plot([datetime.fromisoformat(d) for d, _ in s["curve"]], [e for _, e in s["curve"]], label=f"{s['Risk %']}% risk")
    ax.set_yscale("log")
    ax.set_title("Position sizing sensitivity: Rs 5 lakh, whole lots, modelled option P&L")
    ax.legend()
    ax.grid(alpha=0.3)
    save(fig, "12_position_sizing.png")

    fig, ax = plt.subplots(figsize=(9, 4))
    ax.hist(mc_dd, bins=60, color="#8d6e63")
    ax.axvline(float(dd.max()), color="r", label=f"historical {dd.max():.1f}R")
    ax.axvline(float(np.percentile(mc_dd, 95)), color="k", ls="--", label=f"95th pct {np.percentile(mc_dd, 95):.1f}R")
    ax.set_title("Monte Carlo (trade order shuffled 10,000x): max drawdown in R")
    ax.legend()
    save(fig, "13_monte_carlo_drawdown.png")

    roll = res["rolling"]
    fig, ax = plt.subplots(figsize=(11, 3.5))
    ax.plot([datetime.fromisoformat(d) for d, _ in roll], [v for _, v in roll], color="#3949ab")
    ax.axhline(0, color="k", lw=0.8)
    ax.set_title("Rolling 50-trade expectancy (net R)")
    ax.grid(alpha=0.3)
    save(fig, "14_rolling_expectancy.png")
    return names


# ------------------------------------------------------------------------------------------------------- main


def main() -> None:
    parser = argparse.ArgumentParser(prog="scripts/audit_liquidity.py")
    parser.add_argument("--refresh", action="store_true", help="rebuild the data snapshot (network)")
    args = parser.parse_args()
    settings = get_settings()
    OUT.mkdir(parents=True, exist_ok=True)
    snap = snapshot(args.refresh)
    D = prepare(snap)
    res: dict = {}
    print("data quality ...")
    res["dq"] = data_quality(D, snap)
    print("reproduce production scan ...")
    res["repro"] = reproduce(D)
    assert not res["repro"]["scan_mismatch_days"], res["repro"]["scan_mismatch_days"][:5]

    base = run(D)
    res["baseline"] = stats(base)
    res["baseline_gross"] = stats(base, "R_gross")
    res["baseline_nocool"] = stats(run(D, cool=None))
    res["splits"] = [row(k, v) for k, v in split(base).items()]
    res["by_side"] = [row(s, [t for t in base if t["side"] == s], cols=list(res["baseline"])) for s in ("Long", "Short")]
    res["by_year"] = [row(y, [t for t in base if t["date"][:4] == y]) | {"Year": y} for y in sorted({t["date"][:4] for t in base})]
    res["by_regime"] = [row(g, [t for t in base if t["ctx"]["regime"] == g]) | {"Bucket": g} for g in ("Bull", "Bear", "Sideways")]
    vix_vals = sorted(t["ctx"]["vix"] for t in base)
    lo_v, hi_v = vix_vals[len(vix_vals) // 3], vix_vals[2 * len(vix_vals) // 3]
    res["vix_cuts"] = (lo_v, hi_v)
    res["by_vix"] = [row(n, [t for t in base if f(t["ctx"]["vix"])]) | {"Bucket": n} for n, f in (
        (f"Low <{lo_v:.1f}", lambda v: v < lo_v), (f"Mid {lo_v:.1f}-{hi_v:.1f}", lambda v: lo_v <= v < hi_v), (f"High >={hi_v:.1f}", lambda v: v >= hi_v))]  # fmt: skip
    res["by_time"] = [row(n, [t for t in base if a <= t["time"][11:] < b]) | {"Bucket": n} for n, a, b in (
        ("09:30-11:00", "09:30", "11:00"), ("11:00-13:00", "11:00", "13:00"), ("13:00-14:30", "13:00", "14:31"))]  # fmt: skip
    res["by_expiry_adj"] = [row(n, [t for t in base if f(t["ctx"])]) | {"Bucket": n} for n, f in (
        ("Day before expiry", lambda x: x["pre_expiry"]), ("Day after expiry", lambda x: x["post_expiry"]),
        ("Other days", lambda x: not x["pre_expiry"] and not x["post_expiry"]))]  # fmt: skip
    res["by_exit"] = [{"Exit": k, "Trades": len(v), "Avg R": round(float(np.mean([t["R"] for t in v])), 3)}
                      for k in sorted({t["exit_reason"] for t in base}) for v in [[t for t in base if t["exit_reason"] == k]]]  # fmt: skip
    monthly = defaultdict(float)
    for t in base:
        monthly[t["date"][:7]] += t["R"]
    res["monthly"] = dict(sorted(monthly.items()))
    res["rolling"] = [(base[i]["date"], float(np.mean([t["R"] for t in base[i - 49 : i + 1]]))) for i in range(49, len(base))]
    res["sweep_outcomes"] = sweep_outcomes(D)
    res["depth_buckets"] = []
    for side in ("Long", "Short"):
        for n, a, b in (("1-2", 1, 2), ("2-5", 2, 5), ("5-10", 5, 10), ("10-20", 10, 20), (">20", 20, 1e9)):
            res["depth_buckets"].append(row(f"{side} {n}", [t for t in base if t["side"] == side and a < t["depth"] <= b]))
    print("variants ...")
    V = {}
    V["A Original"] = base
    V["B Both 15m structure"] = run(D, {"short_tf": "15m"})
    V["C Both 5m structure"] = run(D, {"long_tf": "5m"})
    V["D Long 15m / short 5m, no short rejection"] = run(D, {"short_reject": False})
    V["Long reclaim of PDL"] = run(D, {"long_reclaim": True})
    V["Long reclaim + displacement 0.1 ATR"] = run(D, {"long_reclaim": True, "disp": 0.1})
    V["Displacement 0.1 ATR (both)"] = run(D, {"disp": 0.1})
    V["Displacement 0.2 ATR (both)"] = run(D, {"disp": 0.2})
    V["Strict pivots (unique extreme)"] = run(D, {"strict": True})
    for k in (2, 5, 10, 20):
        V[f"Sweep > {k} pts"] = run(D, {"sweep_pts": k})
    for k in (0.01, 0.025, 0.05, 0.1):
        V[f"Sweep > {k} x ATR"] = run(D, {"sweep_atr": k})
    V["VWAP only (no open filter)"] = run(D, {"open_filter": False})
    V["Open only (no VWAP)"] = run(D, {"vwap": False})
    V["Neither VWAP nor open"] = run(D, {"vwap": False, "open_filter": False})
    for m in (0, 20, 40, 50):
        V[f"Min risk {m}"] = run(D, {"min_risk": m})
    V["Sweep-extreme stop"] = run(D, {"stop_at": "sweep"})
    for rm in (1.0, 1.5, 2.5, 3.0):
        V[f"Target {rm}R"] = run(D, exit_kw={"r_mult": rm})
    for mode in ("liquidity", "be", "trail", "partial", "structure"):
        V[f"Exit: {mode}"] = run(D, exit_kw={"mode": mode})
    for et in ("14:45", "15:00", "15:20"):
        V[f"Time exit {et}"] = run(D, exit_kw={"exit_t": et})
    V["Expiry days included"] = run(D, expiry="keep")
    V["Expiry days only"] = run(D, expiry="only")
    V["Also skip day before expiry"] = run(D, day_filter=lambda x, s: not x["pre_expiry"])
    for c in ((0, 0), (2, 3), (3, 1), (3, 5), (4, 3)):
        V[f"Cooldown {c[0]} losses / {c[1]} days" if c[0] else "No cooldown"] = run(D, cool=c if c[0] else None)
    V["HTF trend: with SMA50 only"] = run(D, day_filter=lambda x, s: (x["pdc"] > x["sma50"]) == (s["side"] == "Long"))
    V["HTF trend: against SMA50 only"] = run(D, day_filter=lambda x, s: (x["pdc"] > x["sma50"]) != (s["side"] == "Long"))
    V["Skip gap > 0.5 ATR"] = run(D, day_filter=lambda x, s: abs(s["day_open"] - x["pdc"]) <= 0.5 * x["atr"])
    is_vix = sorted(t["ctx"]["vix"] for t in base if t["date"] <= IS_END)
    is_lo, is_hi = is_vix[len(is_vix) // 3], is_vix[2 * len(is_vix) // 3]
    res["is_vix_cuts"] = (is_lo, is_hi)
    V["Skip high-VIX tercile (IS cut)"] = run(D, day_filter=lambda x, s: x["vix"] < is_hi)
    V["Skip low-VIX tercile (IS cut)"] = run(D, day_filter=lambda x, s: x["vix"] >= is_lo)
    V["Room to opposite level >= 2R"] = run(D, day_filter=lambda x, s: (s["opp"] - s["entry"]) * (1 if s["side"] == "Long" else -1) >= 2 * s["risk"])
    V["Entries before 13:00 only"] = run(D, day_filter=lambda x, s: s["time"][11:] < "13:00")
    V["Dynamic target: opposite level clamped 1-3R"] = run(D, exit_kw={"mode": "dynamic"})
    V["Working tree v3: 3R + stop to entry at 1R"] = run(D, exit_kw={"mode": "be", "r_mult": 3.0})
    res["variants"] = []
    for name, tr in V.items():
        sp = split(tr)
        res["variants"].append(row(name, tr) | {"IS exp R": stats(sp["IS 2016-2020"]).get("Expectancy R"),
                                               "OOS exp R": stats(sp["OOS 2021-2026"]).get("Expectancy R"),
                                               "OOS trades": len(sp["OOS 2021-2026"])})  # fmt: skip
    res["long_only"] = [row(n, [t for t in V[n] if t["side"] == "Long"]) for n in ("A Original", "Long reclaim of PDL", "Long reclaim + displacement 0.1 ATR", "C Both 5m structure")]
    res["short_only"] = [row(n, [t for t in V[n] if t["side"] == "Short"]) for n in ("A Original", "B Both 15m structure", "D Long 15m / short 5m, no short rejection")]
    res["exits_full"] = [row(n, V[n], cols=["Trades", "Win %", "Expectancy R", "PF", "Max DD R", "Avg hold min", "Reached 1R %", "Reached 2R %", "Stopped %", "1R then stopped %"])
                         for n in ("Target 1.0R", "Target 1.5R", "A Original", "Target 2.5R", "Target 3.0R", "Exit: liquidity", "Exit: be", "Exit: trail", "Exit: partial", "Exit: structure", "Time exit 14:45", "Time exit 15:00", "Time exit 15:20")]  # fmt: skip
    res["costs"] = []
    for c in (0, 2, 4, 6, 8, 10):
        rr = [(t["points"] - c) / t["risk"] for t in base]
        res["costs"].append({"Cost index pts/trade": c, "Expectancy R": round(float(np.mean(rr)), 4), "Total R": round(sum(rr), 1),
                             "CI95": "[{:.3f}, {:.3f}]".format(*bootstrap_ci(rr))})  # fmt: skip

    # selection on IS only, judged on OOS once
    base_is = stats(split(base)["IS 2016-2020"])["Expectancy R"]
    cand = [v for v in res["variants"] if v["Config"] in (
        "Long reclaim of PDL", "Displacement 0.1 ATR (both)", "Sweep > 0.025 x ATR", "HTF trend: with SMA50 only",
        "Skip gap > 0.5 ATR", "Skip high-VIX tercile (IS cut)", "Room to opposite level >= 2R", "Entries before 13:00 only",
        "B Both 15m structure", "Strict pivots (unique extreme)") and (v["IS exp R"] or -9) > base_is + 0.02]  # fmt: skip
    res["combo_picked"] = [v["Config"] for v in cand]
    filt, cfg = [], {}
    for v in cand:
        n = v["Config"]
        cfg |= {"Long reclaim of PDL": {"long_reclaim": True}, "Displacement 0.1 ATR (both)": {"disp": 0.1},
                "Sweep > 0.025 x ATR": {"sweep_atr": 0.025}, "B Both 15m structure": {"short_tf": "15m"},
                "Strict pivots (unique extreme)": {"strict": True}}.get(n, {})  # fmt: skip
        f = {"HTF trend: with SMA50 only": lambda x, s: (x["pdc"] > x["sma50"]) == (s["side"] == "Long"),
             "Skip gap > 0.5 ATR": lambda x, s: abs(s["day_open"] - x["pdc"]) <= 0.5 * x["atr"],
             "Skip high-VIX tercile (IS cut)": lambda x, s: x["vix"] < is_hi,
             "Room to opposite level >= 2R": lambda x, s: (s["opp"] - s["entry"]) * (1 if s["side"] == "Long" else -1) >= 2 * s["risk"],
             "Entries before 13:00 only": lambda x, s: s["time"][11:] < "13:00"}.get(n)  # fmt: skip
        if f:
            filt.append(f)
    combo = run(D, cfg, day_filter=lambda x, s: all(f(x, s) for f in filt)) if cand else base
    V["G IS-selected combination"] = combo
    final = {"A Original": base, "B Symmetric structure (both 15m)": V["B Both 15m structure"],
             "C Long PDL reclaim": V["Long reclaim of PDL"], "D Displacement 0.1 ATR": V["Displacement 0.1 ATR (both)"],
             "F ATR sweep > 0.025 ATR": V["Sweep > 0.025 x ATR"], "G IS-selected combination": combo,
             "Working tree v3 (3R + BE at 1R)": V["Working tree v3: 3R + stop to entry at 1R"]}  # fmt: skip
    res["final"] = []
    for name, tr in final.items():
        sp = split(tr)
        res["final"].append(row(name, tr, cols=["Trades", "Trades/yr", "Win %", "Expectancy R", "CI95 low", "CI95 high", "PF", "Max DD R"])
                            | {"IS exp R": stats(sp["IS 2016-2020"]).get("Expectancy R"), "OOS exp R": stats(sp["OOS 2021-2026"]).get("Expectancy R"),
                               "OOS CI95": "[{}, {}]".format(stats(sp["OOS 2021-2026"]).get("CI95 low"), stats(sp["OOS 2021-2026"]).get("CI95 high")),
                               "OOS PF": stats(sp["OOS 2021-2026"]).get("PF")})  # fmt: skip
    print("walk-forward ...")
    res["wf_minrisk"] = walk_forward(D, {"name": "Min risk", "values": [0, 20, 30, 40, 50]}, lambda v: run(D, {"min_risk": v}), base)
    res["wf_target"] = walk_forward(D, {"name": "Target R", "values": [1.0, 1.5, 2.0, 2.5, 3.0]}, lambda v: run(D, exit_kw={"r_mult": v}), base)

    print("monte carlo ...")
    r = np.array([t["R"] for t in base])
    g = np.random.default_rng(SEED)
    shuffled = np.array([g.permutation(r) for _ in range(10000)])
    cum = np.cumsum(shuffled, axis=1)
    mc_dd = (np.maximum.accumulate(np.concatenate([np.zeros((10000, 1)), cum], axis=1), axis=1)[:, 1:] - cum).max(axis=1)
    boot = g.choice(r, (10000, len(r))).sum(axis=1)
    res["mc"] = {"dd_median": float(np.median(mc_dd)), "dd_p95": float(np.percentile(mc_dd, 95)), "dd_p99": float(np.percentile(mc_dd, 99)),
                 "total_p5": float(np.percentile(boot, 5)), "total_median": float(np.median(boot)), "p_total_below_0": float((boot < 0).mean() * 100),
                 "losing_streak_p95": int(np.percentile([max((len(s) for s in "".join("L" if x <= 0 else "W" for x in p).split("W")), default=0) for p in shuffled[:2000]], 95))}  # fmt: skip
    sd = float(np.std(r, ddof=1))
    mu = float(np.mean(r))
    res["sample"] = {"sd_R": sd, "mean_R": mu, "n_needed_95": math.ceil((1.96 * sd / mu) ** 2) if mu > 0 else None,
                     "n_needed_for_0.1R": math.ceil((1.96 * sd / 0.1) ** 2), "tests_run": len(V)}  # fmt: skip

    print("options ...")
    eod = Eod()
    res["expiry_check"] = expiry_check(D, eod)
    res["opt_val"] = option_validation(D, eod)
    opt_rows = []
    for t in base:
        t["opt"] = option_trade(t, settings)
        opt_rows.append(t)
    res["opt_R"] = {"index_R_gross": stats(base, "R_gross")["Expectancy R"], "n": len(base)}
    for slip in (0.0, 0.5, 1.0, 2.0):
        o = [option_trade(t, settings, slip=slip)["opt_R"] for t in base]
        res["opt_R"][f"slip {slip}"] = {"mean": round(float(np.mean(o)), 4), "ci": "[{:.3f}, {:.3f}]".format(*bootstrap_ci(o)), "total": round(sum(o), 1)}
    for m in (0.9, 1.1):
        o = [option_trade(t, settings, iv_exit=m)["opt_R"] for t in base]
        res["opt_R"][f"IV x{m} at exit"] = round(float(np.mean(o)), 4)
    o = [option_trade(t, settings, otm=2)["opt_R"] for t in base]
    res["opt_R"]["2 strikes OTM (working tree), slip 0.5"] = {"mean": round(float(np.mean(o)), 4), "ci": "[{:.3f}, {:.3f}]".format(*bootstrap_ci(o))}
    v3 = V["Working tree v3: 3R + stop to entry at 1R"]
    for otm in (0, 2):
        o = [option_trade(t | {"ctx": t["ctx"]}, settings, otm=otm)["opt_R"] for t in v3]
        res["opt_R"][f"Working tree v3 exits, {otm} OTM, slip 0.5"] = {"mean": round(float(np.mean(o)), 4), "ci": "[{:.3f}, {:.3f}]".format(*bootstrap_ci(o))}
    tgt = [t["opt"]["opt_R"] for t in base if t["exit_reason"] == "Target"]
    stp = [t["opt"]["opt_R"] for t in base if t["exit_reason"] == "Stop"]
    res["opt_R"]["at_target"] = (round(float(np.mean(tgt)), 2), round(float(np.min(tgt)), 2), round(float(np.max(tgt)), 2)) if tgt else None
    res["opt_R"]["at_stop"] = (round(float(np.mean(stp)), 2), round(float(np.min(stp)), 2), round(float(np.max(stp)), 2)) if stp else None
    contract = []
    for t in base:
        if t["date"] < "2020-01-02":
            continue
        i = D["days"].index(t["date"])
        p, real = D["days"][i - 1], eod.nearest(t["date"])
        if not real:
            continue
        kind, k = ("CE" if t["side"] == "Long" else "PE"), round(t["entry"] / 50) * 50
        got = eod.contract_iv(p, real, float(k), kind, D["daily"][p]["close"])
        if not got:
            continue
        ivc, basis = got
        a = option_trade(t, settings, expiry=real)
        b = option_trade(t, settings, iv=ivc, basis=basis, expiry=real)
        contract.append({"vix_R": a["opt_R"], "contract_R": b["opt_R"], "vix_prem": a["premium_in"], "contract_prem": b["premium_in"]})
    res["opt_contract"] = {"n": len(contract), "vix_mean": float(np.mean([c["vix_R"] for c in contract])),
                           "contract_mean": float(np.mean([c["contract_R"] for c in contract])),
                           "prem_ratio_median": float(np.median([c["contract_prem"] / c["vix_prem"] for c in contract]))} if contract else {}  # fmt: skip
    print("sizing ...")
    sizes = [sizing(base, cap, pct) for cap in (100_000, 500_000) for pct in (0.5, 1, 2, 3, 5)]
    res["sizing"] = [{k: v for k, v in s.items() if k != "curve"} for s in sizes]
    o_r = np.array([t["opt"]["opt_R"] for t in base])
    res["ruin"] = [ruin(o_r, pct, len(base)) for pct in (0.5, 1, 2, 3, 5)]
    print("vwap ...")
    res["vwap"] = vwap_study(snap)
    print("charts ...")
    res["charts"] = charts(res, base, opt_rows, sizes, mc_dd)
    write(res)
    workbook(res, base)


def workbook(res: dict, base: list[dict]) -> None:
    """The main RESULTS.md tables plus every baseline trade, as an Excel workbook next to the reports."""
    from openpyxl import (
        Workbook,  # not a project dependency: run with `uv run --with openpyxl`
    )
    from openpyxl.styles import Font, PatternFill

    b, sp = res["baseline"], {r["Config"]: r for r in res["splits"]}
    summary = [
        ["NIFTY PDH/PDL liquidity sweep audit", f"data to {END}; generated by backend/scripts/audit_liquidity.py"],
        ["Verdict", "FAIL: no demonstrated edge after costs; not ready for live capital (see EXECUTIVE_SUMMARY.md)"],
        ["Net R", f"(index points - {COST_PTS:g}) / initial risk; every signal taken, expiry days skipped, 3-loss/3-day cooldown"],
        ["Trades", b["Trades"]], ["Win %", b["Win %"]], ["Expectancy R (net)", b["Expectancy R"]],
        ["CI95", f"[{b['CI95 low']}, {b['CI95 high']}]"], ["Expectancy R (gross)", res["baseline_gross"]["Expectancy R"]],
        ["Profit factor", b["PF"]], ["Max drawdown R", b["Max DD R"]],
        ["IS 2016-2020 expectancy R", sp["IS 2016-2020"]["Expectancy R"]], ["OOS 2021-2026 expectancy R", sp["OOS 2021-2026"]["Expectancy R"]],
        ["Option columns", "MODELLED: Black-76 at previous-day India VIX, forward = spot, IV constant, Rs 0.5 slippage per side. No real intraday option prices exist."],
    ]  # fmt: skip
    trades = [{
        "Date": t["date"], "Side": t["side"], "Entry time": t["time"][11:], "Entry": t["entry"], "Stop": t["stop"],
        "Risk pts": t["risk"], "Sweep depth pts": t["depth"], "Exit time": t["exit_time"][11:], "Exit reason": t["exit_reason"],
        "Exit": t["exit"], "Points": t["points"], "R gross": round(t["R_gross"], 3), "R": round(t["R"], 3),
        "MFE R": round(t["mfe_R"], 2), "MAE R": round(t["mae_R"], 2), "Hold min": t["hold_min"],
        "Option (ATM)": f"{round(t['entry'] / 50) * 50} {'CE' if t['side'] == 'Long' else 'PE'}", "Lot": t["opt"]["lot"],
        "Premium in (model)": round(t["opt"]["premium_in"], 2), "Premium out (model)": round(t["opt"]["premium_out"], 2),
        "Option Rs 1 lot (model)": round(t["opt"]["pnl"]), "Option R (model)": round(t["opt"]["opt_R"], 3),
    } for t in sorted(base, key=lambda t: t["time"])]  # fmt: skip
    sheets = {"Trades": trades, "IS-OOS": res["splits"], "By side": res["by_side"], "By year": res["by_year"],
              "By VIX": res["by_vix"], "By entry time": res["by_time"], "By regime": res["by_regime"],
              "Sweep depth": res["sweep_outcomes"] + res["depth_buckets"], "Variants": res["variants"],
              "Exits": res["exits_full"], "Configs A-G": res["final"], "Costs": res["costs"],
              "Sizing": res["sizing"], "Risk of ruin": res["ruin"], "Option model check": res["opt_val"]["table"],
              "Monthly R": [{"Month": k, "R": round(v, 2)} for k, v in res["monthly"].items()]}  # fmt: skip
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for r in summary:
        ws.append(r)
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 30, 110
    head, red, green = PatternFill("solid", fgColor="1F4E78"), Font(color="C62828"), Font(color="2E7D32")
    for title, rows in sheets.items():
        sheet = wb.create_sheet(title)
        keys = list(dict.fromkeys(k for r in rows for k in r))
        sheet.append(keys)
        for c in sheet[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in rows:
            sheet.append([None if isinstance(v, float) and math.isnan(v) else v if v is None or isinstance(v, (int, float, str))
                          else str(v) for v in (r.get(k) for k in keys)])  # fmt: skip
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = 16
            if col[0].value in ("R", "R gross", "Points", "Option R (model)", "Option Rs 1 lot (model)", "Expectancy R"):
                for c in col[1:]:
                    if isinstance(c.value, (int, float)):
                        c.font = green if c.value > 0 else red if c.value < 0 else Font()
        sheet.freeze_panes, sheet.auto_filter.ref = "A2", sheet.dimensions
    path = OUT / f"NIFTY_Liquidity_Audit_{END}.xlsx"
    wb.save(path)
    print("wrote", path)


def write(res: dict) -> None:
    dq, rp, b = res["dq"], res["repro"], res["baseline"]
    L = [(f"# Liquidity strategy audit: generated results\n\nGenerated by `backend/scripts/audit_liquidity.py` from the frozen snapshot "
         f"`backend/data/intraday/audit_snapshot.json.gz` (data to {END}). Net R = (index points - {COST_PTS:g}) / initial risk. "
         f"IS = {START[:4]}-2020, OOS = 2021-{END[:4]}. Bootstrap CIs: 4,000 resamples, seed {SEED}.\n"),
         "## Data quality\n", f"Sessions {dq['first']} to {dq['last']}, {dq['bars']:,} 5m bars.\n", md(dq["years"]),
         f"\nSessions with fewer than 75 5m bars: {len(dq['short'])}: {dq['short'][:40]}\n\nWeekend sessions: {dq['weekend']}\n",
         "\n5m-derived daily bars vs NSE EOD index bars (2019 - Dec 2025):\n", md(dq["eod_cmp"]),
         f"\nNSE EOD dates missing from the intraday data: {dq['missing_vs_eod']}\n\nIntraday dates absent from NSE EOD: {dq['extra_vs_eod']}\n",
         "\nKaggle vs Groww on the same 5m bars (last 45 days of the Kaggle file):\n", md(dq["groww_cmp"]),
         "\n## Reproduction and look-ahead checks\n",
         (f"Production `scan()` signals on non-expiry days {START}-{END}: {rp['signals']}. Days where `scan_v` defaults differ from `scan()`: "
         f"{len(rp['scan_mismatch_days'])}. Exits where `run_exit` differs from the specified exit (2R, 15:10): {rp['exit_mismatches']}; "
         f"from the working tree's `exit_of()` ({R_MULT:g}R, stop to entry at +1R): {rp['wt_exit_mismatches']}. "
         f"Signals that change when re-scanned with bars only up to the entry bar (repaint): {rp['repaint']}. "
         f"Signals already present one bar earlier: {rp['early']}.\n"),
         "\n## Baseline (original strategy: expiry filter, 3-loss/3-day cooldown, 2R, 15:10)\n", md([b]), "\nGross (no costs):\n", md([res["baseline_gross"]]),
         "\nWithout the cooldown:\n", md([res["baseline_nocool"]]),
         "\n### In-sample / out-of-sample\n", md(res["splits"]), "\n### By side (long = PDL sweep, short = PDH raid)\n", md(res["by_side"]),
         "\n### By year\n", md(res["by_year"], ["Year", *SHORT_COLS]), "\n### By trend regime (previous close vs SMA50 and its 20-day slope)\n", md(res["by_regime"], ["Bucket", *SHORT_COLS]),
         f"\n### By India VIX (previous day, tercile cuts {res['vix_cuts'][0]:.2f} / {res['vix_cuts'][1]:.2f})\n", md(res["by_vix"], ["Bucket", *SHORT_COLS]),
         "\n### By entry time\n", md(res["by_time"], ["Bucket", *SHORT_COLS]), "\n### Expiry-adjacent sessions\n", md(res["by_expiry_adj"], ["Bucket", *SHORT_COLS]),
         "\n### By exit reason\n", md(res["by_exit"]), "\n### Monthly net R\n", md([{"Month": k, "R": round(v, 2)} for k, v in res["monthly"].items()]),
         "\n## Sweep depth\n\nAll sweep days (not only trades): close back inside the level or beyond it.\n", md(res["sweep_outcomes"]),
         "\nBaseline trades by sweep depth at entry:\n", md(res["depth_buckets"]),
         "\n## Variants (one change at a time vs the original)\n", md(res["variants"], ["Config", *SHORT_COLS, "IS exp R", "OOS exp R", "OOS trades"]),
         "\n### Long side only\n", md(res["long_only"]), "\n### Short side only\n", md(res["short_only"]),
         "\n### Exits\n", md(res["exits_full"]), "\n### Cost sensitivity (original)\n", md(res["costs"]),
         f"\n## Final comparison\n\nG = every candidate whose IS expectancy beat the original by > 0.02R, combined: {res['combo_picked']}\n", md(res["final"]),
         "\n## Walk-forward (anchored, yearly re-selection from 2019)\n"]  # fmt: skip
    for k in ("wf_minrisk", "wf_target"):
        w = res[k]
        L += [f"\n### {k}\n", md(w["picked"]), "\n", md([{"Run": "Walk-forward", **{c: w["wf"].get(c) for c in SHORT_COLS}},
                                                         {"Run": "Original (fixed)", **{c: w["base"].get(c) for c in SHORT_COLS}}])]  # fmt: skip
    mc, sm = res["mc"], res["sample"]
    L += ["\n## Monte Carlo and sample size\n",
          (f"Trade order shuffled 10,000x: max drawdown median {mc['dd_median']:.1f}R, 95th pct {mc['dd_p95']:.1f}R, 99th pct {mc['dd_p99']:.1f}R; "
          f"95th pct longest losing streak {mc['losing_streak_p95']}. Bootstrap total R: 5th pct {mc['total_p5']:.1f}, median {mc['total_median']:.1f}, "
          f"P(total < 0) {mc['p_total_below_0']:.1f}%.\n\nPer-trade SD {sm['sd_R']:.3f}R, mean {sm['mean_R']:.4f}R: trades needed for a 95% CI "
          f"excluding zero at this mean: {sm['n_needed_95']}; to resolve +-0.1R: {sm['n_needed_for_0.1R']}. Configurations tested: {sm['tests_run']} "
          f"(Bonferroni 5% => per-test p < {0.05 / sm['tests_run']:.4f}, |t| > {abs(NormalDist().inv_cdf(0.025 / sm['tests_run'])):.2f}).\n"),
          "\n## Options\n", f"\nExpiry calendar: `expiry_for()` vs NSE's nearest listed expiry (2020 on): {len(res['expiry_check'])} mismatching days\n",
          md(res["expiry_check"][:40]),
          (f"\nBlack-76 vs real NSE closes, nearest-expiry ATM, {res['opt_val']['days']} days. Parity forward basis F - spot: median "
          f"{res['opt_val']['basis_median']:.1f} pts (P10 {res['opt_val']['basis_p10']:.1f}, P90 {res['opt_val']['basis_p90']:.1f}). Skew: 2% OTM put IV - ATM IV "
          f"median {(res['opt_val']['skew_put_2pct'] or 0) * 100:.2f} vol pts; 2% OTM call {(res['opt_val']['skew_call_2pct'] or 0) * 100:.2f}.\n"), md(res["opt_val"]["table"]),
          "\nOption R per trade (1 lot ATM, previous-day VIX Black-76, charges, slippage per side):\n",
          md([{"Model": k, "Value": str(v)} for k, v in res["opt_R"].items()]),
          f"\nContract-IV check (2020 on, previous-day IV of the traded strike/expiry + parity forward): {res['opt_contract']}\n",
          "\n## Position sizing (whole lots, compounding, modelled option P&L)\n", md(res["sizing"]),
          "\nRisk of ruin (bootstrap of per-trade option R, fractional sizing):\n", md(res["ruin"]),
          "\n## VWAP study (last 60 days)\n", md([{k: v for k, v in res["vwap"].items() if k not in ("signals", "only_live", "only_twap")}]),
          (f"\nSignals: { {k: len(v) for k, v in res['vwap']['signals'].items()} }\n\nOnly with NIFTYBEES VWAP: {res['vwap']['only_live']}\n\n"
          f"Only with index TWAP: {res['vwap']['only_twap']}\n\nAll NIFTYBEES-VWAP signals: {res['vwap']['signals']['NIFTYBEES VWAP (live)']}\n"),
          "\n## Charts\n", *(f"- ![{c}](charts/{c})\n" for c in res["charts"])]  # fmt: skip
    (OUT / "RESULTS.md").write_text("".join(L))
    clean = {k: v for k, v in res.items() if k not in ("charts",)}
    clean["opt_val"] = {k: v for k, v in res["opt_val"].items() if k != "rows"}
    (OUT / "results.json").write_text(json.dumps(clean, indent=1, default=str))
    print("wrote", OUT / "RESULTS.md")


if __name__ == "__main__":
    main()
